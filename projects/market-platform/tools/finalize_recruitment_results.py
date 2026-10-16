from __future__ import annotations

import argparse
from collections import Counter
from dataclasses import dataclass
from datetime import UTC, datetime
from hashlib import sha256
import json
import math
import os
from pathlib import Path
import re
import shutil
import sys
from tempfile import mkdtemp
from typing import Any, Iterator, Mapping, Sequence, TextIO

from openpyxl import Workbook, load_workbook


SITE_NAMES = ("boss", "job51", "liepin")
SITE_LABELS = {
    "boss": "BOSS直聘",
    "job51": "前程无忧",
    "liepin": "猎聘",
}
RECORD_SITE_NAMES = {
    "boss": frozenset(("boss",)),
    "job51": frozenset(("job51", "51job")),
    "liepin": frozenset(("liepin",)),
}
QUOTA_TARGET = 50
ALLOWED_QUOTA_OUTCOMES = frozenset(("quota_met", "exhausted_below_quota", "empty"))
BLOCKING_QUOTA_OUTCOMES = (
    "pending",
    "running",
    "failed",
    "page_limit_below_quota",
    "deferred",
)
NEGOTIABLE_MARKERS = ("面议", "待定", "另议", "可议", "报价")
NUMBER_RE = re.compile(r"\d+(?:\.\d+)?")
EXCEL_ILLEGAL_RE = re.compile(
    r"[\x00-\x08\x0B\x0C\x0E-\x1F\uD800-\uDFFF\uFFFE\uFFFF]"
)
MAX_STREAM_VALUE_CHARS = 32 * 1024 * 1024

JOB_HEADERS = (
    "记录ID",
    "岗位",
    "薪资",
    "职业标识",
    "分类一级",
    "分类二级",
    "分类三级",
    "公司",
    "地点",
    "岗位链接",
)
RELATION_HEADERS = (
    "记录ID",
    "岗位",
    "薪资",
    "职业标识",
    "分类一级",
    "分类二级",
    "分类三级",
    "分类路径",
    "岗位链接",
)


class FinalizationError(RuntimeError):
    """The persisted crawl does not satisfy the final delivery contract."""


@dataclass(frozen=True, slots=True)
class CategoryPath:
    path_id: str
    level_1: str | None
    level_2: str | None
    level_3: str | None
    path: str


@dataclass(frozen=True, slots=True)
class CategoryUnit:
    key: str
    name: str
    paths: tuple[CategoryPath, ...]
    quota_outcome: str
    expected_records: int


@dataclass(frozen=True, slots=True)
class CoverageContract:
    site: str
    target: int
    numeric_salary_only: bool
    units: Mapping[str, CategoryUnit]
    quota_counts: Mapping[str, int]
    reported_records: int | None


@dataclass(frozen=True, slots=True)
class RecordRelation:
    category_key: str
    category_path: CategoryPath


def _as_nonnegative_int(value: object, *, label: str) -> int:
    if isinstance(value, bool):
        raise FinalizationError(f"{label} must be a non-negative integer")
    try:
        parsed = int(value)
    except (TypeError, ValueError) as exc:
        raise FinalizationError(f"{label} must be a non-negative integer") from exc
    if parsed < 0:
        raise FinalizationError(f"{label} must be a non-negative integer")
    return parsed


def _required_text(value: object, *, label: str) -> str:
    text = str(value or "").strip()
    if not text:
        raise FinalizationError(f"{label} must not be empty")
    return text


def _optional_text(value: object) -> str | None:
    text = str(value or "").strip()
    return text or None


def _load_json_object(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise FinalizationError(f"missing required file: {path}")
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise FinalizationError(f"invalid JSON object: {path}: {exc}") from exc
    if not isinstance(value, dict):
        raise FinalizationError(f"expected a JSON object: {path}")
    return value


def _split_path(raw: object) -> tuple[str | None, str | None, str | None, str]:
    text = _required_text(raw, label="category path")
    parts = [part.strip() for part in text.split(">") if part.strip()]
    level_1 = parts[0] if parts else None
    level_2 = parts[1] if len(parts) > 1 else None
    level_3 = parts[2] if len(parts) > 2 else (parts[-1] if parts else None)
    return level_1, level_2, level_3, text


def _path_from_boss(value: Mapping[str, Any]) -> CategoryPath:
    raw_path = value.get("category_path_raw")
    if raw_path:
        fallback_1, fallback_2, fallback_3, path = _split_path(raw_path)
    else:
        fallback_1 = fallback_2 = fallback_3 = None
        parts = [
            _optional_text(value.get("category_level_1_raw")),
            _optional_text(value.get("category_level_2_raw")),
            _optional_text(value.get("category_level_3_raw")),
        ]
        path = " > ".join(part for part in parts if part)
        if not path:
            raise FinalizationError("BOSS category relation has no path")
    return CategoryPath(
        path_id=path,
        level_1=_optional_text(value.get("category_level_1_raw")) or fallback_1,
        level_2=_optional_text(value.get("category_level_2_raw")) or fallback_2,
        level_3=_optional_text(value.get("category_level_3_raw")) or fallback_3,
        path=path,
    )


def _path_from_job51(value: Mapping[str, Any]) -> CategoryPath:
    level_1_code = _required_text(value.get("level_1_code"), label="51job level_1_code")
    level_2_code = _required_text(value.get("level_2_code"), label="51job level_2_code")
    level_3_code = _required_text(value.get("level_3_code"), label="51job level_3_code")
    level_1 = _required_text(value.get("level_1_name"), label="51job level_1_name")
    level_2 = _required_text(value.get("level_2_name"), label="51job level_2_name")
    level_3 = _required_text(value.get("level_3_name"), label="51job level_3_name")
    expected_path_id = f"{level_1_code}/{level_2_code}/{level_3_code}"
    supplied_path_id = _optional_text(value.get("path_key"))
    if supplied_path_id is not None and supplied_path_id != expected_path_id:
        raise FinalizationError(
            "51job path_key does not match its category code fields: "
            f"{supplied_path_id} != {expected_path_id}"
        )
    return CategoryPath(
        path_id=supplied_path_id or expected_path_id,
        level_1=level_1,
        level_2=level_2,
        level_3=level_3,
        path=" > ".join((level_1, level_2, level_3)),
    )


def _path_from_liepin(value: Mapping[str, Any]) -> CategoryPath:
    path = _required_text(
        value.get("path") or value.get("category_path_raw"),
        label="Liepin category path",
    )
    fallback_1, fallback_2, fallback_3, _ = _split_path(path)
    return CategoryPath(
        path_id=path,
        level_1=_optional_text(value.get("level_1"))
        or _optional_text(value.get("category_level_1_raw"))
        or fallback_1,
        level_2=_optional_text(value.get("level_2"))
        or _optional_text(value.get("category_level_2_raw"))
        or fallback_2,
        level_3=_optional_text(value.get("level_3"))
        or _optional_text(value.get("category_level_3_raw"))
        or fallback_3,
        path=path,
    )


def _inventory_from_coverage(
    site: str, coverage: Mapping[str, Any]
) -> dict[str, tuple[CategoryPath, ...]]:
    grouped: dict[str, dict[str, CategoryPath]] = {}
    if site == "boss":
        rows = coverage.get("positions")
        if not isinstance(rows, list) or not rows:
            raise FinalizationError("boss inventory is missing coverage.positions")
        for raw in rows:
            if not isinstance(raw, Mapping):
                raise FinalizationError("boss inventory position must be an object")
            key = _required_text(raw.get("position_code"), label="boss position_code")
            raw_paths = raw.get("category_paths_raw")
            raw_paths = (
                raw_paths
                if isinstance(raw_paths, list) and raw_paths
                else [raw.get("category_path_raw")]
            )
            for raw_path in raw_paths:
                level_1, level_2, level_3, path = _split_path(raw_path)
                grouped.setdefault(key, {})[path] = CategoryPath(
                    path_id=path,
                    level_1=level_1,
                    level_2=level_2,
                    level_3=level_3,
                    path=path,
                )
    elif site == "job51":
        rows = coverage.get("categories")
        if not isinstance(rows, list) or not rows:
            raise FinalizationError("job51 inventory is missing coverage.categories")
        for raw in rows:
            if not isinstance(raw, Mapping):
                raise FinalizationError("job51 inventory category must be an object")
            key = _required_text(raw.get("search_code"), label="job51 search_code")
            path = _path_from_job51(raw)
            grouped.setdefault(key, {})[path.path_id] = path
    elif site == "liepin":
        rows = coverage.get("categories")
        if not isinstance(rows, list) or not rows:
            raise FinalizationError("liepin inventory is missing coverage.categories")
        for raw in rows:
            if not isinstance(raw, Mapping):
                raise FinalizationError("liepin inventory category must be an object")
            key = _required_text(raw.get("url"), label="liepin category url")
            path = _path_from_liepin(raw)
            if key in grouped:
                raise FinalizationError(f"liepin inventory has duplicate category url: {key}")
            grouped[key] = {path.path_id: path}
    else:
        raise FinalizationError(f"unsupported site: {site}")
    return {
        key: tuple(paths[path_id] for path_id in sorted(paths))
        for key, paths in sorted(grouped.items())
    }


def _validate_quota_counts(
    site: str,
    quota_counts: Mapping[str, Any],
    *,
    total_units: int,
) -> dict[str, int]:
    parsed = {
        name: _as_nonnegative_int(quota_counts.get(name, 0), label=f"{site}.{name}")
        for name in (*ALLOWED_QUOTA_OUTCOMES, *BLOCKING_QUOTA_OUTCOMES)
    }
    blocking = {name: parsed[name] for name in BLOCKING_QUOTA_OUTCOMES if parsed[name]}
    if blocking:
        details = ", ".join(f"{name}={count}" for name, count in blocking.items())
        raise FinalizationError(f"{site} quota has unfinished or invalid states: {details}")
    terminal_total = sum(parsed[name] for name in ALLOWED_QUOTA_OUTCOMES)
    if terminal_total != total_units:
        raise FinalizationError(
            f"{site} quota state total {terminal_total} does not match inventory {total_units}"
        )
    return parsed


def _assert_status_total(
    site: str,
    *,
    expected_total: int,
    values: Mapping[str, object],
) -> None:
    parsed = {
        key: _as_nonnegative_int(value, label=f"{site}.{key}")
        for key, value in values.items()
    }
    actual_total = sum(parsed.values())
    if actual_total != expected_total:
        raise FinalizationError(
            f"{site} status total {actual_total} does not match inventory {expected_total}"
        )


def _boss_contract(coverage: Mapping[str, Any]) -> CoverageContract:
    run_config = coverage.get("run_config")
    run_config = run_config if isinstance(run_config, Mapping) else {}
    target = _as_nonnegative_int(
        run_config.get("per_position_limit")
        or (coverage.get("quota_summary") or {}).get("quota_target"),
        label="boss quota target",
    )
    numeric = bool(run_config.get("numeric_price_only", False))
    inventory = _inventory_from_coverage("boss", coverage)
    declared_units = _as_nonnegative_int(
        coverage.get("position_codes_total"), label="boss position_codes_total"
    )
    if declared_units != len(inventory):
        raise FinalizationError(
            f"boss inventory total {declared_units} does not match "
            f"coverage positions {len(inventory)}"
        )
    _assert_status_total(
        "boss",
        expected_total=declared_units,
        values={
            "completed": coverage.get("position_codes_completed", 0),
            "pending": coverage.get("position_codes_pending", 0),
            "deferred": coverage.get("position_codes_deferred", 0),
            "failed": coverage.get("position_codes_failed", 0),
        },
    )
    if coverage.get("position_paths_total") is not None:
        declared_paths = _as_nonnegative_int(
            coverage.get("position_paths_total"), label="boss position_paths_total"
        )
        actual_paths = sum(len(paths) for paths in inventory.values())
        if declared_paths != actual_paths:
            raise FinalizationError(
                f"boss inventory total paths {declared_paths} does not match {actual_paths}"
            )
    raw_quota = coverage.get("quota_summary")
    if not isinstance(raw_quota, Mapping):
        raise FinalizationError("boss coverage is missing quota_summary")
    merged_quota = dict(raw_quota)
    merged_quota["pending"] = coverage.get("position_codes_pending", 0)
    merged_quota["deferred"] = coverage.get("position_codes_deferred", 0)
    merged_quota["failed"] = max(
        _as_nonnegative_int(raw_quota.get("failed", 0), label="boss.failed"),
        _as_nonnegative_int(
            coverage.get("position_codes_failed", 0),
            label="boss.position_codes_failed",
        ),
    )
    quota = _validate_quota_counts("boss", merged_quota, total_units=len(inventory))
    rows = coverage.get("positions")
    if not isinstance(rows, list):
        raise FinalizationError("boss coverage positions must be a list")
    units: dict[str, CategoryUnit] = {}
    for raw in rows:
        if not isinstance(raw, Mapping):
            raise FinalizationError("boss coverage position must be an object")
        key = _required_text(raw.get("position_code"), label="boss position_code")
        outcome = _required_text(raw.get("quota_status"), label=f"boss {key} quota_status")
        expected = _as_nonnegative_int(raw.get("records_fetched", 0), label=f"boss {key} records")
        units[key] = CategoryUnit(
            key=key,
            name=str(raw.get("position_name") or key),
            paths=inventory[key],
            quota_outcome=outcome,
            expected_records=expected,
        )
    reported = coverage.get("records_fetched")
    return CoverageContract(
        site="boss",
        target=target,
        numeric_salary_only=numeric,
        units=units,
        quota_counts=quota,
        reported_records=(
            _as_nonnegative_int(reported, label="boss records_fetched")
            if reported is not None
            else None
        ),
    )


def _job51_contract(coverage: Mapping[str, Any]) -> CoverageContract:
    raw_quota = coverage.get("quota_counts")
    if not isinstance(raw_quota, Mapping):
        raise FinalizationError("job51 coverage is missing quota_counts")
    target = _as_nonnegative_int(
        raw_quota.get("target_per_unique_search_code"),
        label="job51 quota target",
    )
    numeric = bool(raw_quota.get("numeric_salary_only", False))
    inventory = _inventory_from_coverage("job51", coverage)
    declared_units = _as_nonnegative_int(
        coverage.get("unique_search_codes"), label="job51 unique_search_codes"
    )
    quota_units = _as_nonnegative_int(
        raw_quota.get("unique_search_codes_total"),
        label="job51 unique_search_codes_total",
    )
    if declared_units != len(inventory) or quota_units != len(inventory):
        raise FinalizationError(
            f"job51 inventory total does not match categories: "
            f"coverage={declared_units}, quota={quota_units}, actual={len(inventory)}"
        )
    declared_paths = _as_nonnegative_int(
        coverage.get("official_category_paths"), label="job51 official_category_paths"
    )
    actual_paths = sum(len(paths) for paths in inventory.values())
    if declared_paths != actual_paths:
        raise FinalizationError(
            f"job51 inventory total paths {declared_paths} does not match {actual_paths}"
        )
    raw_status_counts = coverage.get("status_counts")
    if isinstance(raw_status_counts, Mapping):
        _assert_status_total(
            "job51",
            expected_total=declared_paths,
            values={str(key): value for key, value in raw_status_counts.items()},
        )
    quota = _validate_quota_counts("job51", raw_quota, total_units=len(inventory))
    rows = coverage.get("categories")
    if not isinstance(rows, list):
        raise FinalizationError("job51 coverage categories must be a list")
    unit_values: dict[str, tuple[str, int, str]] = {}
    for raw in rows:
        if not isinstance(raw, Mapping):
            raise FinalizationError("job51 coverage category must be an object")
        key = _required_text(raw.get("search_code"), label="job51 search_code")
        outcome = _required_text(raw.get("quota_outcome"), label=f"job51 {key} quota_outcome")
        expected = _as_nonnegative_int(raw.get("items_count", 0), label=f"job51 {key} items_count")
        name = str(raw.get("level_3_name") or key)
        current = (outcome, expected, name)
        previous = unit_values.get(key)
        if previous is not None and previous[:2] != current[:2]:
            raise FinalizationError(f"job51 alias rows disagree for search code {key}")
        unit_values[key] = current
    units = {
        key: CategoryUnit(
            key=key,
            name=unit_values[key][2],
            paths=paths,
            quota_outcome=unit_values[key][0],
            expected_records=unit_values[key][1],
        )
        for key, paths in inventory.items()
    }
    reported = coverage.get("record_count")
    return CoverageContract(
        site="job51",
        target=target,
        numeric_salary_only=numeric,
        units=units,
        quota_counts=quota,
        reported_records=(
            _as_nonnegative_int(reported, label="job51 record_count")
            if reported is not None
            else None
        ),
    )


def _liepin_contract(coverage: Mapping[str, Any]) -> CoverageContract:
    raw_quota = coverage.get("quota_outcomes")
    if not isinstance(raw_quota, Mapping):
        raise FinalizationError("liepin coverage is missing quota_outcomes")
    target = _as_nonnegative_int(coverage.get("quota_target"), label="liepin quota target")
    numeric = bool(coverage.get("numeric_salary_only", False))
    inventory = _inventory_from_coverage("liepin", coverage)
    summary = coverage.get("summary")
    summary = summary if isinstance(summary, Mapping) else {}
    declared_units = _as_nonnegative_int(
        summary.get("categories_total"), label="liepin categories_total"
    )
    if declared_units != len(inventory):
        raise FinalizationError(
            f"liepin inventory total {declared_units} does not match "
            f"coverage categories {len(inventory)}"
        )
    _assert_status_total(
        "liepin",
        expected_total=declared_units,
        values={
            "done": summary.get("categories_done", 0),
            "empty": summary.get("categories_empty", 0),
            "failed": summary.get("categories_failed", 0),
            "pending": summary.get("categories_pending", 0),
            "running": summary.get("categories_running", 0),
        },
    )
    quota = _validate_quota_counts("liepin", raw_quota, total_units=len(inventory))
    rows = coverage.get("categories")
    if not isinstance(rows, list):
        raise FinalizationError("liepin coverage categories must be a list")
    units: dict[str, CategoryUnit] = {}
    for raw in rows:
        if not isinstance(raw, Mapping):
            raise FinalizationError("liepin coverage category must be an object")
        key = _required_text(raw.get("url"), label="liepin category url")
        outcome = _required_text(raw.get("quota_outcome"), label=f"liepin {key} quota_outcome")
        expected = _as_nonnegative_int(
            raw.get("record_count", 0), label=f"liepin {key} record_count"
        )
        units[key] = CategoryUnit(
            key=key,
            name=str(raw.get("level_3") or key),
            paths=inventory[key],
            quota_outcome=outcome,
            expected_records=expected,
        )
    reported = summary.get("records_total")
    return CoverageContract(
        site="liepin",
        target=target,
        numeric_salary_only=numeric,
        units=units,
        quota_counts=quota,
        reported_records=(
            _as_nonnegative_int(reported, label="liepin records_total")
            if reported is not None
            else None
        ),
    )


def _coverage_contract(site: str, coverage: Mapping[str, Any]) -> CoverageContract:
    contract = {
        "boss": _boss_contract,
        "job51": _job51_contract,
        "liepin": _liepin_contract,
    }[site](coverage)
    if contract.target != QUOTA_TARGET:
        raise FinalizationError(
            f"{site} quota target is {contract.target}; expected {QUOTA_TARGET}"
        )
    if not contract.numeric_salary_only:
        raise FinalizationError(f"{site} numeric salary mode is not enabled")
    for unit in contract.units.values():
        if unit.quota_outcome not in ALLOWED_QUOTA_OUTCOMES:
            raise FinalizationError(
                f"{site} category {unit.key} has invalid quota outcome {unit.quota_outcome}"
            )
        if unit.expected_records > QUOTA_TARGET:
            raise FinalizationError(
                f"{site} category {unit.key} has {unit.expected_records} records; maximum is 50"
            )
        if unit.quota_outcome == "quota_met" and unit.expected_records != QUOTA_TARGET:
            raise FinalizationError(
                f"{site} category {unit.key} is quota_met with {unit.expected_records} records"
            )
        if unit.quota_outcome == "empty" and unit.expected_records != 0:
            raise FinalizationError(
                f"{site} category {unit.key} is empty with {unit.expected_records} records"
            )
        if unit.quota_outcome == "exhausted_below_quota" and unit.expected_records >= QUOTA_TARGET:
            raise FinalizationError(
                f"{site} category {unit.key} is exhausted with {unit.expected_records} records"
            )
    declared_outcomes = {
        outcome: int(contract.quota_counts.get(outcome, 0))
        for outcome in ("quota_met", "exhausted_below_quota", "empty")
    }
    unit_outcomes = Counter(unit.quota_outcome for unit in contract.units.values())
    actual_outcomes = {
        outcome: unit_outcomes.get(outcome, 0)
        for outcome in ("quota_met", "exhausted_below_quota", "empty")
    }
    if declared_outcomes != actual_outcomes:
        raise FinalizationError(
            f"{site} quota counts do not match unique category outcomes: "
            f"declared={declared_outcomes}, actual={actual_outcomes}"
        )
    return contract


def _path_identity(path: CategoryPath) -> tuple[str, str | None, str | None, str | None, str]:
    return (
        path.path_id,
        path.level_1,
        path.level_2,
        path.level_3,
        path.path,
    )


def _inventory_signature(inventory: Mapping[str, Sequence[CategoryPath]]) -> str:
    payload = [
        [key, [_path_identity(path) for path in inventory[key]]]
        for key in sorted(inventory)
    ]
    encoded = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    return sha256(encoded.encode("utf-8")).hexdigest()


def _assert_inventory_matches(
    site: str,
    actual: Mapping[str, Sequence[CategoryPath]],
    probe: Mapping[str, Sequence[CategoryPath]],
) -> None:
    actual_map = {key: {_path_identity(path) for path in paths} for key, paths in actual.items()}
    probe_map = {key: {_path_identity(path) for path in paths} for key, paths in probe.items()}
    if actual_map == probe_map:
        return
    missing_units = sorted(set(probe_map) - set(actual_map))[:10]
    extra_units = sorted(set(actual_map) - set(probe_map))[:10]
    changed_units = sorted(
        key
        for key in set(actual_map) & set(probe_map)
        if actual_map[key] != probe_map[key]
    )[:10]
    raise FinalizationError(
        f"{site} inventory differs from per1 probe: "
        f"missing={missing_units}, extra={extra_units}, changed={changed_units}"
    )


def _validate_probe_per_category(
    site: str,
    coverage: Mapping[str, Any],
) -> str:
    """Prove that a supplied inventory scope was collected with a limit of one."""

    explicit: list[tuple[str, object]] = []
    if site == "boss":
        run_config = coverage.get("run_config")
        if isinstance(run_config, Mapping) and run_config.get("per_position_limit") is not None:
            explicit.append(
                ("run_config.per_position_limit", run_config["per_position_limit"])
            )
        quota_summary = coverage.get("quota_summary")
        if isinstance(quota_summary, Mapping) and quota_summary.get("quota_target") is not None:
            explicit.append(("quota_summary.quota_target", quota_summary["quota_target"]))
    elif site == "job51":
        run_scope = coverage.get("run_scope")
        if isinstance(run_scope, Mapping) and run_scope.get("per_category_limit") is not None:
            explicit.append(("run_scope.per_category_limit", run_scope["per_category_limit"]))
        quota_counts = coverage.get("quota_counts")
        if (
            isinstance(quota_counts, Mapping)
            and quota_counts.get("target_per_unique_search_code") is not None
        ):
            explicit.append(
                (
                    "quota_counts.target_per_unique_search_code",
                    quota_counts["target_per_unique_search_code"],
                )
            )
    elif site == "liepin" and coverage.get("quota_target") is not None:
        explicit.append(("quota_target", coverage["quota_target"]))

    rows_key = "positions" if site == "boss" else "categories"
    rows = coverage.get(rows_key)
    if isinstance(rows, list):
        row_targets = {
            _as_nonnegative_int(
                row.get("quota_target"),
                label=f"{site} probe {rows_key}[].quota_target",
            )
            for row in rows
            if isinstance(row, Mapping) and row.get("quota_target") is not None
        }
        if row_targets:
            if len(row_targets) != 1:
                raise FinalizationError(
                    f"{site} probe has inconsistent per-category targets: {sorted(row_targets)}"
                )
            explicit.append((f"{rows_key}[].quota_target", next(iter(row_targets))))

    if explicit:
        invalid = [
            (name, _as_nonnegative_int(value, label=f"{site} probe {name}"))
            for name, value in explicit
            if _as_nonnegative_int(value, label=f"{site} probe {name}") != 1
        ]
        if invalid:
            details = ", ".join(f"{name}={value}" for name, value in invalid)
            raise FinalizationError(
                f"{site} probe per-category target must be 1: {details}"
            )
        return ", ".join(name for name, _value in explicit)

    if site != "liepin":
        raise FinalizationError(
            f"{site} probe has no explicit per-category target; expected 1"
        )
    if not isinstance(rows, list) or not rows:
        raise FinalizationError("liepin probe has no categories for per1 validation")
    for index, row in enumerate(rows):
        if not isinstance(row, Mapping):
            raise FinalizationError("liepin probe category must be an object")
        count = _as_nonnegative_int(
            row.get("record_count"),
            label=f"liepin probe categories[{index}].record_count",
        )
        if count > 1:
            raise FinalizationError(
                "liepin legacy probe cannot prove per-category target 1: "
                f"categories[{index}].record_count={count}"
            )
    return "categories[].record_count<=1 (legacy coverage)"


class _JsonStream:
    """Incrementally decode top-level JSON values without materializing the file."""

    def __init__(self, handle: TextIO, *, chunk_size: int = 64 * 1024) -> None:
        self.handle = handle
        self.chunk_size = chunk_size
        self.buffer = ""
        self.position = 0
        self.eof = False
        self.decoder = json.JSONDecoder()

    def _compact(self) -> None:
        if self.position and (
            self.position >= self.chunk_size or self.position > len(self.buffer) // 2
        ):
            self.buffer = self.buffer[self.position :]
            self.position = 0

    def _fill(self) -> bool:
        if self.eof:
            return False
        self._compact()
        chunk = self.handle.read(self.chunk_size)
        if chunk:
            self.buffer += chunk
            return True
        self.eof = True
        return False

    def _skip_space(self) -> None:
        while True:
            while self.position < len(self.buffer) and self.buffer[self.position].isspace():
                self.position += 1
            if self.position < len(self.buffer) or not self._fill():
                return

    def peek(self) -> str:
        self._skip_space()
        if self.position >= len(self.buffer):
            raise FinalizationError("unexpected end of JSON stream")
        return self.buffer[self.position]

    def consume(self, expected: str) -> None:
        actual = self.peek()
        if actual != expected:
            raise FinalizationError(
                f"invalid JSON stream: expected {expected!r}, found {actual!r}"
            )
        self.position += 1

    def value(self) -> object:
        self._skip_space()
        while True:
            try:
                value, end = self.decoder.raw_decode(self.buffer, self.position)
            except json.JSONDecodeError as exc:
                buffered = len(self.buffer) - self.position
                if buffered > MAX_STREAM_VALUE_CHARS:
                    raise FinalizationError(
                        "one JSON value exceeds the 32 MiB streaming safety limit"
                    ) from exc
                if self._fill():
                    continue
                raise FinalizationError(f"invalid JSON stream: {exc}") from exc
            self.position = end
            self._compact()
            return value


def _iter_boss_jobs(path: Path) -> Iterator[dict[str, Any]]:
    if not path.is_file():
        raise FinalizationError(f"missing required file: {path}")
    try:
        handle = path.open("r", encoding="utf-8", newline="")
    except OSError as exc:
        raise FinalizationError(f"cannot open {path}: {exc}") from exc
    with handle:
        stream = _JsonStream(handle)
        stream.consume("{")
        found = False
        first = True
        while stream.peek() != "}":
            if not first:
                stream.consume(",")
            first = False
            key = stream.value()
            if not isinstance(key, str):
                raise FinalizationError("BOSS jobs JSON has a non-string object key")
            stream.consume(":")
            if key != "jobs":
                stream.value()
                continue
            if found:
                raise FinalizationError("BOSS jobs JSON contains duplicate jobs keys")
            found = True
            stream.consume("[")
            array_first = True
            while stream.peek() != "]":
                if not array_first:
                    stream.consume(",")
                array_first = False
                value = stream.value()
                if not isinstance(value, dict):
                    raise FinalizationError("BOSS jobs array must contain objects")
                yield value
            stream.consume("]")
        stream.consume("}")
        if not found:
            raise FinalizationError("BOSS jobs JSON is missing jobs array")


def _iter_jsonl(path: Path) -> Iterator[dict[str, Any]]:
    if not path.is_file():
        raise FinalizationError(f"missing required file: {path}")
    try:
        handle = path.open("r", encoding="utf-8", newline="")
    except OSError as exc:
        raise FinalizationError(f"cannot open {path}: {exc}") from exc
    with handle:
        for line_number, line in enumerate(handle, 1):
            if not line.strip():
                continue
            try:
                value = json.loads(line)
            except json.JSONDecodeError as exc:
                raise FinalizationError(
                    f"invalid JSONL at {path}:{line_number}: {exc}"
                ) from exc
            if not isinstance(value, dict):
                raise FinalizationError(f"JSONL row at {path}:{line_number} is not an object")
            yield value


def _record_path(site: str, scope: Path) -> Path:
    return scope / ("jobs.json" if site == "boss" else "jobs.jsonl")


def _iter_records(site: str, scope: Path) -> Iterator[dict[str, Any]]:
    path = _record_path(site, scope)
    return _iter_boss_jobs(path) if site == "boss" else _iter_jsonl(path)


def _salary_text(record: Mapping[str, Any]) -> str:
    values: list[str] = []
    evidence = record.get("price_evidence")
    if isinstance(evidence, list):
        for item in evidence:
            if not isinstance(item, Mapping):
                continue
            value = str(item.get("raw_text") or "").strip()
            if value and value not in values:
                values.append(value)
    if not values:
        fallback = str(record.get("budget_raw") or record.get("salary") or "").strip()
        if fallback:
            values.append(fallback)
    return " | ".join(values)


def _has_numeric_salary(site: str, record: Mapping[str, Any], salary: str) -> bool:
    if not salary or any(marker in salary for marker in NEGOTIABLE_MARKERS):
        return False
    metadata = record.get("metadata")
    metadata = metadata if isinstance(metadata, Mapping) else {}
    if site == "job51":
        for key in ("salary_min", "salary_max"):
            value = metadata.get(key)
            if (
                isinstance(value, (int, float))
                and not isinstance(value, bool)
                and math.isfinite(float(value))
                and float(value) > 0
            ):
                return True
    salary_component = re.split(r"[·•]", salary, maxsplit=1)[0]
    return any(float(value) > 0 for value in NUMBER_RE.findall(salary_component))


def _boss_record_relations(record: Mapping[str, Any]) -> list[RecordRelation]:
    matches = record.get("category_matches")
    if not isinstance(matches, list) or not matches:
        matches = [record]
    relations: list[RecordRelation] = []
    for match in matches:
        if not isinstance(match, Mapping):
            raise FinalizationError("BOSS category match must be an object")
        key = _required_text(
            match.get("position_code") or record.get("position_code"),
            label="BOSS record position_code",
        )
        paths = match.get("category_paths")
        if isinstance(paths, list) and paths:
            relation_paths = [
                _path_from_boss(path)
                for path in paths
                if isinstance(path, Mapping)
            ]
            if len(relation_paths) != len(paths):
                raise FinalizationError("BOSS category_paths must contain objects")
        else:
            raw_paths = match.get("category_paths_raw")
            if isinstance(raw_paths, list) and raw_paths:
                relation_paths = []
                for raw_path in raw_paths:
                    level_1, level_2, level_3, path = _split_path(raw_path)
                    relation_paths.append(
                        CategoryPath(path, level_1, level_2, level_3, path)
                    )
            else:
                relation_paths = [_path_from_boss(match)]
        relations.extend(RecordRelation(key, path) for path in relation_paths)
    return relations


def _job51_record_relations(record: Mapping[str, Any]) -> list[RecordRelation]:
    metadata = record.get("metadata")
    metadata = metadata if isinstance(metadata, Mapping) else {}
    paths = record.get("category_paths")
    if not isinstance(paths, list) or not paths:
        paths = metadata.get("category_paths")
    if not isinstance(paths, list) or not paths:
        raise FinalizationError("job51 record has no category_paths")
    relations: list[RecordRelation] = []
    for raw in paths:
        if not isinstance(raw, Mapping):
            raise FinalizationError("job51 category_paths must contain objects")
        path = _path_from_job51(raw)
        key = _required_text(raw.get("level_3_code"), label="job51 record search_code")
        relations.append(RecordRelation(key, path))
    return relations


def _liepin_record_relations(record: Mapping[str, Any]) -> list[RecordRelation]:
    key = _required_text(
        record.get("category_source_url") or record.get("list_url"),
        label="liepin record category URL",
    )
    return [RecordRelation(key, _path_from_liepin(record))]


def _record_relations(site: str, record: Mapping[str, Any]) -> list[RecordRelation]:
    relations = {
        "boss": _boss_record_relations,
        "job51": _job51_record_relations,
        "liepin": _liepin_record_relations,
    }[site](record)
    if not relations:
        raise FinalizationError(f"{site} record has no category relation")
    seen: set[tuple[str, str]] = set()
    for relation in relations:
        key = (relation.category_key, relation.category_path.path_id)
        if key in seen:
            raise FinalizationError(
                f"{site} record has duplicate category relation: {key[0]} / {key[1]}"
            )
        seen.add(key)
    return relations


def _validate_record(
    site: str,
    record: Mapping[str, Any],
    contract: CoverageContract,
) -> tuple[str, str, list[RecordRelation]]:
    record_id = _required_text(record.get("record_id"), label=f"{site} record_id")
    record_site = _required_text(record.get("site_name"), label=f"{site} site_name")
    if record_site not in RECORD_SITE_NAMES[site]:
        raise FinalizationError(
            f"record {record_id} belongs to {record_site}, expected {site}"
        )
    _required_text(record.get("title_raw"), label=f"{site} record {record_id} title")
    if record.get("category_identity_verified") is not True:
        raise FinalizationError(f"{site} record {record_id} category is not verified")
    salary = _salary_text(record)
    if not _has_numeric_salary(site, record, salary):
        raise FinalizationError(f"{site} record {record_id} has no numeric salary")
    relations = _record_relations(site, record)
    relations_by_category: dict[str, dict[str, CategoryPath]] = {}
    for relation in relations:
        unit = contract.units.get(relation.category_key)
        if unit is None:
            raise FinalizationError(
                f"{site} record {record_id} references unknown category {relation.category_key}"
            )
        known_paths = {path.path_id for path in unit.paths}
        if relation.category_path.path_id not in known_paths:
            raise FinalizationError(
                f"{site} record {record_id} references unknown inventory path "
                f"{relation.category_path.path_id}"
            )
        relations_by_category.setdefault(relation.category_key, {})[
            relation.category_path.path_id
        ] = relation.category_path
    if site in ("boss", "job51"):
        for category_key, actual_paths in relations_by_category.items():
            expected_paths = {
                path.path_id: path for path in contract.units[category_key].paths
            }
            missing = sorted(set(expected_paths) - set(actual_paths))
            extra = sorted(set(actual_paths) - set(expected_paths))
            if missing or extra:
                raise FinalizationError(
                    f"{site} record {record_id} does not contain all official category aliases "
                    f"for {category_key}: missing={missing}, extra={extra}"
                )
            changed = sorted(
                path_id
                for path_id in expected_paths
                if _path_identity(actual_paths[path_id])
                != _path_identity(expected_paths[path_id])
            )
            if changed:
                raise FinalizationError(
                    f"{site} record {record_id} category path fields differ from inventory "
                    f"for {category_key}: changed={changed}"
                )
    return record_id, salary, relations


def _validate_site_records(
    site: str,
    scope: Path,
    contract: CoverageContract,
) -> dict[str, int]:
    seen_record_ids: set[str] = set()
    associations = {key: set() for key in contract.units}
    expanded_relations = 0
    for record in _iter_records(site, scope):
        record_id, _salary, relations = _validate_record(site, record, contract)
        if record_id in seen_record_ids:
            raise FinalizationError(f"{site} contains duplicate record_id {record_id}")
        seen_record_ids.add(record_id)
        expanded_relations += len(relations)
        for category_key in {relation.category_key for relation in relations}:
            bucket = associations[category_key]
            if record_id in bucket:
                raise FinalizationError(
                    f"{site} duplicate category association: {category_key} / {record_id}"
                )
            bucket.add(record_id)
            if len(bucket) > QUOTA_TARGET:
                raise FinalizationError(
                    f"{site} category {category_key} exceeds 50 records"
                )
    for key, unit in contract.units.items():
        actual = len(associations[key])
        if actual != unit.expected_records:
            raise FinalizationError(
                f"{site} category {key} has {actual} record associations; "
                f"coverage reports {unit.expected_records}"
            )
    if contract.reported_records is not None and len(seen_record_ids) != contract.reported_records:
        raise FinalizationError(
            f"{site} has {len(seen_record_ids)} records; coverage reports "
            f"{contract.reported_records}"
        )
    return {
        "records": len(seen_record_ids),
        "category_associations": sum(len(value) for value in associations.values()),
        "expanded_path_relations": expanded_relations,
    }


def _excel_value(value: object) -> object:
    if not isinstance(value, str):
        return value
    return EXCEL_ILLEGAL_RE.sub("", value)[:32767]


def _record_link(record: Mapping[str, Any]) -> str | None:
    return _optional_text(record.get("detail_url") or record.get("list_url"))


def _job_row(
    record: Mapping[str, Any],
    *,
    record_id: str,
    salary: str,
    primary: RecordRelation,
) -> tuple[object, ...]:
    path = primary.category_path
    return tuple(
        _excel_value(value)
        for value in (
            record_id,
            str(record.get("title_raw") or ""),
            salary,
            primary.category_key,
            path.level_1,
            path.level_2,
            path.level_3,
            _optional_text(record.get("publisher_name_raw") or record.get("company")),
            _optional_text(record.get("location_raw") or record.get("location")),
            _record_link(record),
        )
    )


def _relation_row(
    record: Mapping[str, Any],
    *,
    record_id: str,
    salary: str,
    relation: RecordRelation,
) -> tuple[object, ...]:
    path = relation.category_path
    return tuple(
        _excel_value(value)
        for value in (
            record_id,
            str(record.get("title_raw") or ""),
            salary,
            relation.category_key,
            path.level_1,
            path.level_2,
            path.level_3,
            path.path,
            _record_link(record),
        )
    )


def _save_workbook(workbook: Workbook, path: Path) -> None:
    workbook.save(path)
    with path.open("rb+") as handle:
        os.fsync(handle.fileno())


def _verify_workbook_row_count(
    path: Path,
    *,
    sheet_name: str,
    expected_rows: int,
) -> None:
    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        if sheet_name not in workbook.sheetnames:
            raise FinalizationError(
                f"workbook is missing required sheet {sheet_name!r}: {path}"
            )
        actual_rows = sum(
            1 for _row in workbook[sheet_name].iter_rows(values_only=True)
        )
    except FinalizationError:
        raise
    except Exception as exc:
        raise FinalizationError(f"cannot verify saved workbook {path}: {exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()
    if actual_rows != expected_rows:
        raise FinalizationError(
            f"workbook row count mismatch for {path}: "
            f"expected={expected_rows}, actual={actual_rows}"
        )


def _export_site_workbooks(
    site: str,
    scope: Path,
    contract: CoverageContract,
    output_dir: Path,
    *,
    expected_records: int,
    expected_relations: int,
) -> None:
    output_dir.mkdir(parents=True, exist_ok=False)
    jobs_book = Workbook(write_only=True)
    jobs_sheet = jobs_book.create_sheet("岗位")
    jobs_sheet.append(JOB_HEADERS)
    relations_book = Workbook(write_only=True)
    relations_sheet = relations_book.create_sheet("岗位职业关联")
    relations_sheet.append(RELATION_HEADERS)
    for record in _iter_records(site, scope):
        record_id, salary, relations = _validate_record(site, record, contract)
        jobs_sheet.append(
            _job_row(
                record,
                record_id=record_id,
                salary=salary,
                primary=relations[0],
            )
        )
        for relation in relations:
            relations_sheet.append(
                _relation_row(
                    record,
                    record_id=record_id,
                    salary=salary,
                    relation=relation,
                )
            )
    jobs_path = output_dir / "jobs.xlsx"
    relations_path = output_dir / "job_category_relations.xlsx"
    _save_workbook(jobs_book, jobs_path)
    _save_workbook(relations_book, relations_path)
    _verify_workbook_row_count(
        jobs_path,
        sheet_name="岗位",
        expected_rows=expected_records + 1,
    )
    _verify_workbook_row_count(
        relations_path,
        sheet_name="岗位职业关联",
        expected_rows=expected_relations + 1,
    )


def _source_fingerprint(path: Path) -> tuple[int, int]:
    try:
        stat = path.stat()
    except OSError as exc:
        raise FinalizationError(f"cannot stat required source file {path}: {exc}") from exc
    if not path.is_file():
        raise FinalizationError(f"required source is not a file: {path}")
    return stat.st_size, stat.st_mtime_ns


def _validate_scope_mapping(
    scopes: Mapping[str, Path], *, label: str
) -> dict[str, Path]:
    unknown = sorted(set(scopes) - set(SITE_NAMES))
    missing = sorted(set(SITE_NAMES) - set(scopes))
    if unknown or missing:
        raise FinalizationError(
            f"{label} must contain exactly boss, job51 and liepin; "
            f"missing={missing}, unknown={unknown}"
        )
    resolved: dict[str, Path] = {}
    for site in SITE_NAMES:
        path = Path(scopes[site]).expanduser().resolve()
        if not path.is_dir():
            raise FinalizationError(f"{label} scope is not a directory: {path}")
        resolved[site] = path
    return resolved


def _atomic_write_json(path: Path, value: Mapping[str, Any]) -> None:
    temporary = path.with_name(f".{path.name}.tmp")
    with temporary.open("w", encoding="utf-8", newline="\n") as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
        handle.flush()
        os.fsync(handle.fileno())
    os.replace(temporary, path)


def finalize_recruitment_results(
    *,
    per50_scopes: Mapping[str, Path],
    output_dir: Path,
    per1_scopes: Mapping[str, Path] | None = None,
) -> dict[str, Any]:
    scopes = _validate_scope_mapping(per50_scopes, label="per50_scopes")
    probes = (
        _validate_scope_mapping(per1_scopes, label="per1_scopes")
        if per1_scopes is not None
        else None
    )
    output_dir = Path(output_dir).expanduser().resolve()
    if output_dir.exists():
        raise FinalizationError(f"output directory already exists: {output_dir}")
    output_dir.parent.mkdir(parents=True, exist_ok=True)
    temporary_dir = Path(
        mkdtemp(prefix=f".{output_dir.name}.tmp-", dir=output_dir.parent)
    )
    try:
        report: dict[str, Any] = {
            "valid": True,
            "generated_at": datetime.now(UTC).isoformat(timespec="seconds"),
            "quota_target": QUOTA_TARGET,
            "numeric_salary_only": True,
            "inventory_probe_supplied": probes is not None,
            "sites": {},
        }
        totals = {
            "records": 0,
            "category_associations": 0,
            "expanded_path_relations": 0,
            "inventory_units": 0,
            "inventory_paths": 0,
        }
        for site in SITE_NAMES:
            coverage_path = scopes[site] / "coverage.json"
            coverage_before = _source_fingerprint(coverage_path)
            coverage = _load_json_object(coverage_path)
            if _source_fingerprint(coverage_path) != coverage_before:
                raise FinalizationError(f"{site} coverage changed while it was read")
            contract = _coverage_contract(site, coverage)
            inventory = {key: unit.paths for key, unit in contract.units.items()}
            probe_matches: bool | None = None
            probe_signature: str | None = None
            probe_per_category_validation: str | None = None
            if probes is not None:
                probe_path = probes[site] / "coverage.json"
                probe_before = _source_fingerprint(probe_path)
                probe_coverage = _load_json_object(probe_path)
                if _source_fingerprint(probe_path) != probe_before:
                    raise FinalizationError(
                        f"{site} probe coverage changed while it was read"
                    )
                probe_per_category_validation = _validate_probe_per_category(
                    site, probe_coverage
                )
                probe_inventory = _inventory_from_coverage(site, probe_coverage)
                _assert_inventory_matches(site, inventory, probe_inventory)
                probe_matches = True
                probe_signature = _inventory_signature(probe_inventory)

            source_path = _record_path(site, scopes[site])
            before = _source_fingerprint(source_path)
            metrics = _validate_site_records(site, scopes[site], contract)
            if _source_fingerprint(source_path) != before:
                raise FinalizationError(f"{site} records changed during validation")
            _export_site_workbooks(
                site,
                scopes[site],
                contract,
                temporary_dir / site,
                expected_records=metrics["records"],
                expected_relations=metrics["expanded_path_relations"],
            )
            if _source_fingerprint(source_path) != before:
                raise FinalizationError(f"{site} records changed during export")
            if _source_fingerprint(coverage_path) != coverage_before:
                raise FinalizationError(f"{site} coverage changed during export")

            quota_counts = {
                key: int(contract.quota_counts.get(key, 0))
                for key in ("quota_met", "exhausted_below_quota", "empty")
            }
            site_report = {
                "site_label": SITE_LABELS[site],
                **metrics,
                "inventory_units": len(inventory),
                "inventory_paths": sum(len(paths) for paths in inventory.values()),
                "inventory_signature": _inventory_signature(inventory),
                "inventory_matches_probe": probe_matches,
                "probe_inventory_signature": probe_signature,
                "probe_per_category_validation": probe_per_category_validation,
                "quota": quota_counts,
                "source_scope": os.fspath(scopes[site]),
                "outputs": {
                    "jobs": f"{site}/jobs.xlsx",
                    "job_category_relations": f"{site}/job_category_relations.xlsx",
                },
            }
            report["sites"][site] = site_report
            for key in totals:
                totals[key] += int(site_report[key])
        report["totals"] = totals
        _atomic_write_json(temporary_dir / "validation.json", report)
        try:
            os.replace(temporary_dir, output_dir)
        except OSError as exc:
            raise FinalizationError(
                f"could not atomically publish {output_dir}: {exc}"
            ) from exc
        return report
    except Exception:
        if temporary_dir.exists():
            shutil.rmtree(temporary_dir)
        raise


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Validate and stream-package three completed recruitment per50 scopes."
    )
    for site in SITE_NAMES:
        parser.add_argument(f"--{site}-scope", required=True, type=Path)
        parser.add_argument(f"--{site}-probe-scope", type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = _build_parser().parse_args(argv)
    per50 = {site: getattr(args, f"{site}_scope") for site in SITE_NAMES}
    probe_values = {
        site: getattr(args, f"{site}_probe_scope") for site in SITE_NAMES
    }
    supplied = [value is not None for value in probe_values.values()]
    if any(supplied) and not all(supplied):
        print(
            json.dumps(
                {
                    "valid": False,
                    "error": "all three --*-probe-scope arguments are required together",
                },
                ensure_ascii=False,
                indent=2,
            ),
            file=sys.stderr,
        )
        return 2
    try:
        report = finalize_recruitment_results(
            per50_scopes=per50,
            per1_scopes=probe_values if all(supplied) else None,
            output_dir=args.output_dir,
        )
    except (FinalizationError, OSError) as exc:
        print(
            json.dumps(
                {"valid": False, "error": str(exc)},
                ensure_ascii=False,
                indent=2,
            ),
            file=sys.stderr,
        )
        return 1
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
