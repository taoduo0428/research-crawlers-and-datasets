from __future__ import annotations

import argparse
from collections import defaultdict
from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from hashlib import sha256
import json
import os
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


SOURCE_NAME = "任务与价格.xlsx"
OUTPUT_NAME = "分类价格统计.xlsx"
SOURCE_SHEET = "任务与价格"
SOURCE_HEADERS = ["任务/岗位", "分类一级", "分类二级", "分类三级", "价格"]
SITES = ("猪八戒", "一品威客", "实现网")
LEVEL_NAMES = ("一级", "二级", "三级")
MAX_SOURCE_ROWS = 20_000
MONEY_QUANTUM = Decimal("0.01")


class CategoryStatisticsError(RuntimeError):
    pass


@dataclass(frozen=True, slots=True)
class PriceRow:
    categories: tuple[str, ...]
    price: Decimal


@dataclass(frozen=True, slots=True)
class CategoryStatistic:
    categories: tuple[str, ...]
    count: int
    average: Decimal
    median: Decimal


def _sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _read_source_rows(path: Path) -> list[PriceRow]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except (OSError, ValueError) as exc:
        raise CategoryStatisticsError(f"cannot read source workbook: {path}") from exc
    try:
        if SOURCE_SHEET not in workbook.sheetnames:
            raise CategoryStatisticsError(f"missing source sheet in {path}")
        sheet = workbook[SOURCE_SHEET]
        headers = [cell.value for cell in sheet[1]]
        if headers != SOURCE_HEADERS:
            raise CategoryStatisticsError(f"unexpected source headers in {path}: {headers}")
        if sheet.max_row <= 1 or sheet.max_row - 1 > MAX_SOURCE_ROWS:
            raise CategoryStatisticsError(
                f"invalid source row count in {path}: {sheet.max_row - 1}"
            )

        rows: list[PriceRow] = []
        depths: set[int] = set()
        for row_number, cells in enumerate(
            sheet.iter_rows(min_row=2, max_col=5), start=2
        ):
            if any(cell.data_type == "f" for cell in cells):
                raise CategoryStatisticsError(
                    f"formula is not allowed in source row {row_number}: {path}"
                )
            title = str(cells[0].value or "").strip()
            raw_levels = tuple(str(cell.value or "").strip() for cell in cells[1:4])
            categories = tuple(level for level in raw_levels if level)
            if not title or not categories:
                raise CategoryStatisticsError(
                    f"missing title or category in source row {row_number}: {path}"
                )
            if raw_levels[: len(categories)] != categories:
                raise CategoryStatisticsError(
                    f"non-contiguous category levels in source row {row_number}: {path}"
                )
            price_value = cells[4].value
            if (
                isinstance(price_value, bool)
                or not isinstance(price_value, (int, float))
                or price_value <= 0
            ):
                raise CategoryStatisticsError(
                    f"price is not a positive number in source row {row_number}: {path}"
                )
            rows.append(PriceRow(categories, Decimal(str(price_value))))
            depths.add(len(categories))
    finally:
        workbook.close()

    if len(depths) != 1:
        raise CategoryStatisticsError(
            f"mixed category depths are not supported in {path}: {sorted(depths)}"
        )
    return rows


def _median(values: list[Decimal]) -> Decimal:
    ordered = sorted(values)
    middle = len(ordered) // 2
    if len(ordered) % 2:
        return ordered[middle]
    return (ordered[middle - 1] + ordered[middle]) / 2


def _calculate_statistics(
    rows: list[PriceRow], depth: int
) -> list[CategoryStatistic]:
    groups: dict[tuple[str, ...], list[Decimal]] = defaultdict(list)
    for row in rows:
        if len(row.categories) < depth:
            raise CategoryStatisticsError(
                f"category depth {depth} is unavailable for {row.categories}"
            )
        groups[row.categories[:depth]].append(row.price)

    statistics = [
        CategoryStatistic(
            categories=categories,
            count=len(prices),
            average=sum(prices, Decimal(0)) / len(prices),
            median=_median(prices),
        )
        for categories, prices in groups.items()
    ]
    return sorted(
        statistics,
        key=lambda item: (-item.count, tuple(value.casefold() for value in item.categories)),
    )


def _excel_number(value: Decimal) -> int | float:
    rounded = value.quantize(MONEY_QUANTUM, rounding=ROUND_HALF_UP)
    integral = rounded.to_integral_value()
    return int(integral) if rounded == integral else float(rounded)


def _write_statistics_workbook(
    path: Path, statistics_by_depth: dict[int, list[CategoryStatistic]]
) -> None:
    workbook = Workbook()
    default_sheet = workbook.active
    header_fill = PatternFill("solid", fgColor="1F4E78")

    for index, (depth, statistics) in enumerate(statistics_by_depth.items()):
        sheet = default_sheet if index == 0 else workbook.create_sheet()
        sheet.title = f"{LEVEL_NAMES[depth - 1]}分类统计"
        headers = [f"分类{level}" for level in LEVEL_NAMES[:depth]]
        headers.extend(["任务数量", "平均价格", "中位价格"])
        sheet.append(headers)
        for statistic in statistics:
            sheet.append(
                [
                    *statistic.categories,
                    statistic.count,
                    _excel_number(statistic.average),
                    _excel_number(statistic.median),
                ]
            )

        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column_index in range(1, depth + 1):
            sheet.column_dimensions[
                sheet.cell(row=1, column=column_index).column_letter
            ].width = 24
        count_column = depth + 1
        average_column = depth + 2
        median_column = depth + 3
        sheet.column_dimensions[
            sheet.cell(row=1, column=count_column).column_letter
        ].width = 12
        for column_index in (average_column, median_column):
            sheet.column_dimensions[
                sheet.cell(row=1, column=column_index).column_letter
            ].width = 16
        for row in sheet.iter_rows(min_row=2):
            row[count_column - 1].number_format = "0"
            row[average_column - 1].number_format = "0.00"
            row[median_column - 1].number_format = "0.00"
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    path.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(suffix=".xlsx", delete=False, dir=path.parent) as handle:
        temporary = Path(handle.name)
    try:
        workbook.save(temporary)
        workbook.close()
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def _expected_sheet_rows(
    depth: int, statistics: list[CategoryStatistic]
) -> list[tuple[Any, ...]]:
    headers = tuple(
        [f"分类{level}" for level in LEVEL_NAMES[:depth]]
        + ["任务数量", "平均价格", "中位价格"]
    )
    values = [
        (
            *item.categories,
            item.count,
            _excel_number(item.average),
            _excel_number(item.median),
        )
        for item in statistics
    ]
    return [headers, *values]


def _validate_statistics_workbook(
    path: Path,
    source_count: int,
    statistics_by_depth: dict[int, list[CategoryStatistic]],
) -> None:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        expected_names = [
            f"{LEVEL_NAMES[depth - 1]}分类统计" for depth in statistics_by_depth
        ]
        if workbook.sheetnames != expected_names:
            raise CategoryStatisticsError(
                f"unexpected statistic sheets in {path}: {workbook.sheetnames}"
            )
        for depth, statistics in statistics_by_depth.items():
            sheet = workbook[f"{LEVEL_NAMES[depth - 1]}分类统计"]
            actual = list(sheet.iter_rows(values_only=True))
            expected = _expected_sheet_rows(depth, statistics)
            if actual != expected:
                raise CategoryStatisticsError(
                    f"statistic values do not match source data: {path} sheet {sheet.title}"
                )
            count_column = depth
            if sum(int(row[count_column]) for row in actual[1:]) != source_count:
                raise CategoryStatisticsError(
                    f"statistic count does not reconcile: {path} sheet {sheet.title}"
                )
            if any(
                cell.data_type == "f"
                for row in sheet.iter_rows(min_row=2)
                for cell in row
            ):
                raise CategoryStatisticsError(f"formula found in statistic output: {path}")
    finally:
        workbook.close()


def generate_category_statistics(results_dir: Path) -> dict[str, Any]:
    summary: dict[str, Any] = {}
    for site_name in SITES:
        source = results_dir / site_name / SOURCE_NAME
        destination = results_dir / site_name / OUTPUT_NAME
        source_hash = _sha256(source)
        rows = _read_source_rows(source)
        max_depth = len(rows[0].categories)
        statistics_by_depth = {
            depth: _calculate_statistics(rows, depth)
            for depth in range(1, max_depth + 1)
        }
        _write_statistics_workbook(destination, statistics_by_depth)
        _validate_statistics_workbook(destination, len(rows), statistics_by_depth)
        if _sha256(source) != source_hash:
            raise CategoryStatisticsError(f"source workbook changed: {source}")
        summary[site_name] = {
            "source_rows": len(rows),
            "levels": {
                str(depth): len(statistics)
                for depth, statistics in statistics_by_depth.items()
            },
            "output": str(destination),
            "source_unchanged": True,
        }
    return summary


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Create separate category price statistics workbooks."
    )
    parser.add_argument(
        "--results-dir", type=Path, default=Path("outputs/外包--最终结果")
    )
    args = parser.parse_args()
    try:
        summary = generate_category_statistics(args.results_dir.resolve())
    except (OSError, ValueError, CategoryStatisticsError) as exc:
        print(f"ERROR: {exc}")
        return 1
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
