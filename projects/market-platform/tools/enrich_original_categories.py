from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
import re
import shutil
from tempfile import TemporaryDirectory
from typing import Any
from urllib.parse import urlparse
from uuid import uuid4

from openpyxl import load_workbook

from drv_crawler.categories import (
    CategoryParseError,
    OriginalCategory,
    category_from_zbj_list_url,
    parse_epwk_detail_category,
    parse_shixian_detail_category,
    validate_original_category_fields,
)
from drv_crawler.core import CrawlerError, ProjectConfig, ResilientHttpClient
from drv_crawler.exporters import CSV_FIELDS, write_csv, write_jsonl, write_xlsx
from drv_crawler.models import CrawlRecord, PriceEvidence


MAX_RECORDS_PER_SITE = 20_000
CATEGORY_FIELDS = (
    "category_level_1_raw",
    "category_level_2_raw",
    "category_level_3_raw",
    "category_path_raw",
    "category_source",
    "category_source_url",
    "category_identity_verified",
)
SITE_RECORD_PATHS = {
    "zbj": Path("outputs/外包--完整数据/猪八戒/zbj/records.jsonl"),
    "epwk": Path("outputs/外包--完整数据/一品威客/epwk/records.jsonl"),
    "shixian": Path("outputs/外包--完整数据/实现网/records.jsonl"),
}
VERIFIED_CATEGORY_RECOVERIES = {
    "epwk": {
        "epwk-c3847a26f0a6237969b6": {
            "record_id": "epwk-c3847a26f0a6237969b6",
            "detail_url": "https://task.epwk.com/1156205/",
            "title_raw": "校园欺凌视频素材查找和剪辑",
            "category_level_1_raw": "视频",
            "category_level_2_raw": "影视后期",
            "category_level_3_raw": "剪辑服务",
            "category_path_raw": "视频 > 影视后期 > 剪辑服务",
            "category_source": (
                "search_indexed_original_category_listing+original_site_taxonomy"
            ),
            "category_source_url": "https://task.epwk.com/ypjj/",
            "category_identity_verified": True,
            "category_evidence": {
                "detail_status": 404,
                "identity_method": "exact_title_match_in_search_index",
                "indexed_result_url": "https://task.epwk.com/videomake/",
                "indexed_result_label": "剪辑服务",
                "indexed_result_text": "【剪辑服务】校园欺凌视频素材查找和剪辑",
                "taxonomy_url": "https://talent.epwk.com/ypjj/",
                "task_category_url": "https://task.epwk.com/ypjj/",
                "verified_on": "2026-07-22",
            },
        }
    }
}


class CategoryEnrichmentError(RuntimeError):
    pass


def _read_records(
    path: Path, *, allow_empty: bool = False
) -> list[dict[str, Any]]:
    try:
        rows = [
            json.loads(line)
            for line in path.read_text(encoding="utf-8").splitlines()
            if line.strip()
        ]
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise CategoryEnrichmentError(f"cannot read records: {path}") from exc
    if (not rows and not allow_empty) or len(rows) > MAX_RECORDS_PER_SITE:
        raise CategoryEnrichmentError(f"invalid record count in {path}: {len(rows)}")
    record_ids = [str(row.get("record_id") or "") for row in rows]
    if not all(record_ids) or len(record_ids) != len(set(record_ids)):
        raise CategoryEnrichmentError(f"missing or duplicate record_id in {path}")
    return rows


def _atomic_write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{uuid4().hex}.tmp")
    try:
        temporary.write_text(
            "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows),
            encoding="utf-8",
        )
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def _same_page(left: str, right: str) -> bool:
    left_url = urlparse(left)
    right_url = urlparse(right)
    return (
        left_url.netloc.lower() == right_url.netloc.lower()
        and left_url.path.rstrip("/") == right_url.path.rstrip("/")
    )


def _category_checkpoint_row(
    record: dict[str, Any], category: OriginalCategory
) -> dict[str, Any]:
    return {
        "record_id": str(record["record_id"]),
        "detail_url": str(record.get("detail_url") or ""),
        "title_raw": str(record.get("title_raw") or ""),
        **category.record_fields(),
    }


def _load_checkpoint(path: Path) -> dict[str, dict[str, Any]]:
    if not path.exists():
        return {}
    rows = _read_records(path, allow_empty=True)
    output: dict[str, dict[str, Any]] = {}
    for row in rows:
        record_id = str(row["record_id"])
        if not row.get("category_path_raw") or not row.get("category_identity_verified"):
            raise CategoryEnrichmentError(f"invalid category checkpoint row: {record_id}")
        output[record_id] = row
    return output


def collect_remote_categories(
    *,
    site_name: str,
    records_path: Path,
    config_path: Path,
    state_dir: Path,
    partition_count: int,
    partition_index: int,
    limit: int | None,
) -> dict[str, Any]:
    if site_name not in {"epwk", "shixian"}:
        raise CategoryEnrichmentError("remote collection supports epwk or shixian")
    if partition_count < 1 or partition_count > 8:
        raise CategoryEnrichmentError("partition_count must be between 1 and 8")
    if partition_index < 0 or partition_index >= partition_count:
        raise CategoryEnrichmentError("partition_index is outside partition_count")

    all_records = _read_records(records_path)
    records = [
        row for index, row in enumerate(all_records) if index % partition_count == partition_index
    ]
    if limit is not None:
        records = records[:limit]
    checkpoint_path = (
        state_dir / f"{site_name}-part-{partition_index}-of-{partition_count}.jsonl"
    )
    failure_path = (
        state_dir / f"{site_name}-part-{partition_index}-of-{partition_count}-failures.jsonl"
    )
    completed = _load_checkpoint(checkpoint_path)
    expected_ids = {str(row["record_id"]) for row in records}
    completed = {
        record_id: row for record_id, row in completed.items() if record_id in expected_ids
    }

    config = ProjectConfig(config_path)
    client = ResilientHttpClient(config, site_name=site_name)
    if site_name == "epwk":
        state_path = Path(
            str(config.site("epwk").get("storage_state") or ".crawler-state/epwk-storage-state.json")
        )
        client.load_storage_state(
            state_path, allowed_domains={"epwk.com", "task.epwk.com"}
        )

    failures: list[dict[str, Any]] = []
    consecutive_failures = 0
    pending = [row for row in records if str(row["record_id"]) not in completed]
    try:
        for position, record in enumerate(pending, start=1):
            record_id = str(record["record_id"])
            detail_url = str(record.get("detail_url") or "")
            try:
                if not detail_url:
                    raise CategoryParseError("detail URL missing")
                page = client.get(detail_url, referer=str(record.get("list_url") or ""))
                if not _same_page(page.final_url, detail_url):
                    raise CategoryParseError(
                        f"final URL identity mismatch: {page.final_url}"
                    )
                soup = page.soup()
                if site_name == "epwk":
                    category = parse_epwk_detail_category(
                        soup,
                        detail_url=detail_url,
                        expected_titles=tuple(
                            str(value)
                            for value in (
                                record.get("title_raw"),
                                record.get("detail_title_raw"),
                            )
                            if value
                        ),
                    )
                else:
                    category = parse_shixian_detail_category(
                        soup,
                        detail_url=detail_url,
                        expected_title=str(record.get("title_raw") or ""),
                    )
                completed[record_id] = _category_checkpoint_row(record, category)
                consecutive_failures = 0
            except (CrawlerError, CategoryParseError, OSError, UnicodeError) as exc:
                failures.append(
                    {
                        "record_id": record_id,
                        "detail_url": detail_url,
                        "error_type": type(exc).__name__,
                        "message": str(exc),
                    }
                )
                consecutive_failures += 1
                if consecutive_failures >= 5:
                    failures.extend(
                        {
                            "record_id": str(remaining["record_id"]),
                            "detail_url": str(remaining.get("detail_url") or ""),
                            "error_type": "NotAttemptedAfterFailureStreak",
                            "message": "stopped after five consecutive failures",
                        }
                        for remaining in pending[position:]
                    )
                    break
            if position % 10 == 0 or position == len(pending):
                ordered = [
                    completed[str(row["record_id"])]
                    for row in records
                    if str(row["record_id"]) in completed
                ]
                _atomic_write_jsonl(checkpoint_path, ordered)
                _atomic_write_jsonl(failure_path, failures)
            if position % 50 == 0:
                print(
                    json.dumps(
                        {
                            "site": site_name,
                            "partition": f"{partition_index}/{partition_count}",
                            "completed": len(completed),
                            "expected": len(records),
                            "failures": len(failures),
                        },
                        ensure_ascii=False,
                    ),
                    flush=True,
                )
    finally:
        client.close()

    ordered = [
        completed[str(row["record_id"])]
        for row in records
        if str(row["record_id"]) in completed
    ]
    _atomic_write_jsonl(checkpoint_path, ordered)
    _atomic_write_jsonl(failure_path, failures)
    return {
        "site": site_name,
        "partition_count": partition_count,
        "partition_index": partition_index,
        "expected": len(records),
        "completed": len(ordered),
        "failures": len(failures),
        "complete": len(ordered) == len(records),
        "checkpoint": str(checkpoint_path),
        "failure_file": str(failure_path),
    }


def _load_site_category_checkpoints(
    state_dir: Path, site_name: str
) -> dict[str, dict[str, Any]]:
    checkpoint_name = re.compile(
        rf"{re.escape(site_name)}-part-(\d+)-of-(\d+)\.jsonl"
    )
    paths = sorted(
        path
        for path in state_dir.glob(f"{site_name}-part-*-of-*.jsonl")
        if checkpoint_name.fullmatch(path.name)
    )
    if not paths:
        raise CategoryEnrichmentError(f"no category checkpoints for {site_name}")
    output: dict[str, dict[str, Any]] = {}
    for path in paths:
        for record_id, row in _load_checkpoint(path).items():
            if record_id in output:
                raise CategoryEnrichmentError(
                    f"duplicate category checkpoint record: {record_id}"
                )
            output[record_id] = row
    return output


def _rebuild_record(row: dict[str, Any]) -> CrawlRecord:
    value = dict(row)
    value.pop("price_raw_joined", None)
    value.pop("core_complete", None)
    raw_evidence = value.pop("price_evidence", [])
    evidence = [PriceEvidence(**item) for item in raw_evidence]
    return CrawlRecord(**value, price_evidence=evidence)


def _validate_staged_exports(directory: Path, expected_count: int) -> None:
    jsonl_rows = _read_records(directory / "records.jsonl")
    if len(jsonl_rows) != expected_count:
        raise CategoryEnrichmentError("staged JSONL record count mismatch")
    if any(
        not row.get("category_path_raw")
        or not row.get("category_source")
        or not row.get("category_identity_verified")
        for row in jsonl_rows
    ):
        raise CategoryEnrichmentError("staged JSONL contains missing categories")

    workbook = load_workbook(directory / "records.xlsx", read_only=True, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    xlsx_count = sheet.max_row - 1
    workbook.close()
    if xlsx_count != expected_count or headers != CSV_FIELDS:
        raise CategoryEnrichmentError("staged XLSX validation failed")


def _replace_site_exports(site_dir: Path, records: list[CrawlRecord]) -> None:
    names = ("records.jsonl", "records.csv", "records.xlsx")
    with TemporaryDirectory(prefix=".category-stage-", dir=site_dir) as temporary_name:
        stage = Path(temporary_name)
        write_jsonl(stage / "records.jsonl", records)
        write_csv(stage / "records.csv", records)
        write_xlsx(stage / "records.xlsx", records)
        _validate_staged_exports(stage, len(records))

        backup = site_dir / f".category-backup-{uuid4().hex}"
        backup.mkdir()
        moved_originals: list[str] = []
        replaced: list[str] = []
        delete_backup = False
        try:
            for name in names:
                os.replace(site_dir / name, backup / name)
                moved_originals.append(name)
            for name in names:
                os.replace(stage / name, site_dir / name)
                replaced.append(name)
            delete_backup = True
        except Exception as exc:
            for name in replaced:
                try:
                    (site_dir / name).unlink(missing_ok=True)
                except OSError:
                    pass
            for name in moved_originals:
                if (backup / name).exists():
                    try:
                        os.replace(backup / name, site_dir / name)
                    except OSError:
                        pass
            unrecovered = [name for name in moved_originals if (backup / name).exists()]
            if unrecovered:
                raise CategoryEnrichmentError(
                    "category export replacement and rollback failed; "
                    f"backup retained at {backup}: {', '.join(unrecovered)}"
                ) from exc
            delete_backup = True
            raise
        finally:
            if delete_backup:
                shutil.rmtree(backup, ignore_errors=True)


def apply_categories(
    *, project_root: Path, state_dir: Path
) -> dict[str, Any]:
    resolved_paths = {
        site: project_root / path for site, path in SITE_RECORD_PATHS.items()
    }
    source_rows = {site: _read_records(path) for site, path in resolved_paths.items()}
    category_rows = {
        "epwk": _load_site_category_checkpoints(state_dir, "epwk"),
        "shixian": _load_site_category_checkpoints(state_dir, "shixian"),
    }
    for site_name, recoveries in VERIFIED_CATEGORY_RECOVERIES.items():
        for record_id, recovery in recoveries.items():
            category_rows[site_name].setdefault(record_id, recovery)

    summaries: dict[str, Any] = {}
    rebuilt_by_site: dict[str, list[CrawlRecord]] = {}
    for site_name, rows in source_rows.items():
        enriched: list[CrawlRecord] = []
        title_equals_path = 0
        title_equals_level = 0
        recovered_category_evidence = 0
        for row in rows:
            record_id = str(row["record_id"])
            category_evidence: dict[str, Any] | None = None
            if site_name == "zbj":
                category = category_from_zbj_list_url(str(row.get("list_url") or ""))
                fields = category.record_fields()
            else:
                checkpoint = category_rows[site_name].get(record_id)
                if checkpoint is None:
                    raise CategoryEnrichmentError(
                        f"missing {site_name} category for record {record_id}"
                    )
                if (
                    checkpoint.get("detail_url") != row.get("detail_url")
                    or checkpoint.get("title_raw") != row.get("title_raw")
                ):
                    raise CategoryEnrichmentError(
                        f"category identity mismatch for record {record_id}"
                    )
                fields = {name: checkpoint.get(name) for name in CATEGORY_FIELDS}
                raw_category_evidence = checkpoint.get("category_evidence")
                if isinstance(raw_category_evidence, dict):
                    category_evidence = dict(raw_category_evidence)
                    recovered_category_evidence += 1

            validate_original_category_fields(
                site_name,
                fields,
                detail_url=str(row.get("detail_url") or "") or None,
            )

            value = dict(row)
            value.update(fields)
            metadata = dict(value.get("metadata") or {})
            original_category_metadata: dict[str, Any] = {
                name: fields.get(name) for name in CATEGORY_FIELDS
            }
            if category_evidence is not None:
                original_category_metadata["evidence"] = category_evidence
            metadata["original_category"] = original_category_metadata
            value["metadata"] = metadata
            title = str(value.get("title_raw") or "").strip()
            category_path = str(fields.get("category_path_raw") or "").strip()
            levels = {
                str(fields.get(name) or "").strip()
                for name in (
                    "category_level_1_raw",
                    "category_level_2_raw",
                    "category_level_3_raw",
                )
                if fields.get(name)
            }
            title_equals_path += title == category_path
            title_equals_level += title in levels
            enriched.append(_rebuild_record(value))

        rebuilt_by_site[site_name] = enriched
        summaries[site_name] = {
            "records": len(enriched),
            "title_equals_category_path": title_equals_path,
            "title_equals_any_category_level": title_equals_level,
            "recovered_category_evidence": recovered_category_evidence,
        }

    for site_name, records in rebuilt_by_site.items():
        _replace_site_exports(resolved_paths[site_name].parent, records)
    return summaries


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Collect and apply original-site task/job categories without title inference."
    )
    parser.add_argument("--config", type=Path, default=Path("crawler.config.yaml"))
    parser.add_argument("--project-root", type=Path, default=Path("."))
    parser.add_argument(
        "--state-dir", type=Path, default=Path(".crawler-state/categories")
    )
    subparsers = parser.add_subparsers(dest="command", required=True)

    collect = subparsers.add_parser("collect")
    collect.add_argument("--site", choices=["epwk", "shixian"], required=True)
    collect.add_argument("--partition-count", type=int, default=1)
    collect.add_argument("--partition-index", type=int, default=0)
    collect.add_argument("--limit", type=int)

    subparsers.add_parser("apply")
    args = parser.parse_args()
    try:
        project_root = args.project_root.resolve()
        state_dir = (project_root / args.state_dir).resolve()
        if args.command == "collect":
            records_path = project_root / SITE_RECORD_PATHS[args.site]
            summary = collect_remote_categories(
                site_name=args.site,
                records_path=records_path,
                config_path=project_root / args.config,
                state_dir=state_dir,
                partition_count=args.partition_count,
                partition_index=args.partition_index,
                limit=args.limit,
            )
            print(json.dumps(summary, ensure_ascii=False, indent=2))
            return 0 if summary["complete"] else 2
        summary = apply_categories(project_root=project_root, state_dir=state_dir)
        print(json.dumps(summary, ensure_ascii=False, indent=2))
        return 0
    except (OSError, CategoryEnrichmentError, CategoryParseError, ValueError) as exc:
        print(f"ERROR: {exc}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
