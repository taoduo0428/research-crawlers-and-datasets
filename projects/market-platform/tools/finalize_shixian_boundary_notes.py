from __future__ import annotations

import argparse
from copy import deepcopy
import csv
import json
import os
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook


MAX_RECORDS = 10_000
BOUNDARY_ASSUMPTION = "site_three_month_label"
STRICT_SCOPE = "strict_three_month_window"
BOUNDARY_SCOPE = "three_month_boundary_uncertain"


class BoundaryNoteError(RuntimeError):
    pass


def _read_jsonl(path: Path) -> list[dict[str, Any]]:
    try:
        rows = [
            json.loads(line)
            for line in path.read_text(encoding="utf-8").splitlines()
            if line.strip()
        ]
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise BoundaryNoteError(f"cannot read {path}") from exc
    if not rows or len(rows) > MAX_RECORDS:
        raise BoundaryNoteError("record count is empty or exceeds the safety limit")
    if not all(isinstance(row, dict) for row in rows):
        raise BoundaryNoteError("records.jsonl contains a non-object row")
    return rows


def _atomic_text(path: Path, text: str) -> None:
    temporary = path.with_name(f".{path.name}.{uuid4().hex}.tmp")
    temporary.write_text(text, encoding="utf-8", newline="")
    os.replace(temporary, path)


def _publication_columns(row: dict[str, Any]) -> tuple[str, str, str]:
    metadata = row.get("metadata")
    publication = metadata.get("publication_time") if isinstance(metadata, dict) else None
    if not isinstance(publication, dict):
        raise BoundaryNoteError(f"{row.get('record_id')} has no publication_time")

    if publication.get("boundary_assumption") == BOUNDARY_ASSUMPTION:
        interval_start = str(publication.get("interval_start") or "")
        interval_end = str(publication.get("interval_end") or "")
        if not interval_start or not interval_end:
            raise BoundaryNoteError(
                f"{row.get('record_id')} has an incomplete boundary interval"
            )
        interval = f"{interval_start} ~ {interval_end}"
        note = (
            f"边界不确定：页面仅显示“3个月前”，平台时间桶约为 {interval_start} "
            f"至 {interval_end}；为避免漏单纳入工作清单，但不作为严格三个月内记录。"
        )
        return BOUNDARY_SCOPE, interval, note

    if publication.get("status") != "include":
        raise BoundaryNoteError(
            f"{row.get('record_id')} is neither strict-in-window nor boundary-included"
        )
    return (
        STRICT_SCOPE,
        "",
        "严格三个月内：页面时间证据可判定为截止日当天或之后。",
    )


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            rows = list(reader)
            headers = list(reader.fieldnames or [])
    except (OSError, UnicodeError, csv.Error) as exc:
        raise BoundaryNoteError(f"cannot read {path}") from exc
    if "record_id" not in headers:
        raise BoundaryNoteError("records.csv has no record_id column")
    return headers, rows


def _write_csv(
    path: Path,
    headers: list[str],
    rows: list[dict[str, str]],
    annotations: dict[str, tuple[str, str, str]],
) -> None:
    extra = [
        "publication_scope",
        "publication_boundary_interval",
        "publication_boundary_note",
    ]
    temporary = path.with_name(f".{path.name}.{uuid4().hex}.tmp")
    with temporary.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers + extra)
        writer.writeheader()
        for row in rows:
            record_id = str(row.get("record_id") or "")
            scope, interval, note = annotations[record_id]
            writer.writerow(
                {
                    **row,
                    "publication_scope": scope,
                    "publication_boundary_interval": interval,
                    "publication_boundary_note": note,
                }
            )
    os.replace(temporary, path)


def _write_xlsx(
    source: Path,
    destination: Path,
    annotations: dict[str, tuple[str, str, str]],
) -> None:
    try:
        workbook = load_workbook(source)
    except (OSError, ValueError) as exc:
        raise BoundaryNoteError(f"cannot read {source}") from exc
    sheet = workbook.active
    headers = {
        str(cell.value): index
        for index, cell in enumerate(sheet[1], start=1)
        if cell.value is not None
    }
    record_id_column = headers.get("record_id")
    if record_id_column is None:
        workbook.close()
        raise BoundaryNoteError("records.xlsx has no record_id column")

    first_extra = sheet.max_column + 1
    extra_headers = [
        "publication_scope",
        "publication_boundary_interval",
        "publication_boundary_note",
    ]
    for offset, name in enumerate(extra_headers):
        sheet.cell(row=1, column=first_extra + offset, value=name)
    seen: set[str] = set()
    for row_index in range(2, sheet.max_row + 1):
        record_id = str(sheet.cell(row=row_index, column=record_id_column).value or "")
        if record_id not in annotations:
            workbook.close()
            raise BoundaryNoteError(f"records.xlsx has unknown record_id {record_id}")
        seen.add(record_id)
        for offset, value in enumerate(annotations[record_id]):
            sheet.cell(row=row_index, column=first_extra + offset, value=value)
    if seen != set(annotations):
        workbook.close()
        raise BoundaryNoteError("records.xlsx is missing one or more record IDs")

    temporary = destination.with_name(f".{destination.name}.{uuid4().hex}.tmp")
    workbook.save(temporary)
    workbook.close()
    os.replace(temporary, destination)


def finalize(
    source_run: Path,
    output_dir: Path,
    *,
    expected_total: int | None = None,
    expected_boundary: int | None = None,
) -> dict[str, Any]:
    source = source_run / "shixian"
    required = [source / "records.jsonl", source / "records.csv", source / "records.xlsx"]
    missing = [str(path) for path in required if not path.is_file()]
    if missing:
        raise BoundaryNoteError(f"missing input artifacts: {missing}")
    if output_dir.exists():
        raise BoundaryNoteError(f"output already exists: {output_dir}")

    rows = _read_jsonl(source / "records.jsonl")
    record_ids = [str(row.get("record_id") or "") for row in rows]
    if not all(record_ids) or len(set(record_ids)) != len(record_ids):
        raise BoundaryNoteError("records.jsonl has missing or duplicate record IDs")
    annotations = {
        record_id: _publication_columns(row)
        for record_id, row in zip(record_ids, rows, strict=True)
    }
    boundary_count = sum(
        scope == BOUNDARY_SCOPE for scope, _, _ in annotations.values()
    )
    strict_count = len(rows) - boundary_count
    if expected_total is not None and len(rows) != expected_total:
        raise BoundaryNoteError(
            f"expected {expected_total} records, found {len(rows)}"
        )
    if expected_boundary is not None and boundary_count != expected_boundary:
        raise BoundaryNoteError(
            f"expected {expected_boundary} boundary records, found {boundary_count}"
        )

    csv_headers, csv_rows = _read_csv(source / "records.csv")
    csv_ids = [str(row.get("record_id") or "") for row in csv_rows]
    if csv_ids != record_ids:
        raise BoundaryNoteError("CSV and JSONL record order/IDs do not match")

    output_dir.mkdir(parents=True)
    enriched_rows: list[dict[str, Any]] = []
    for row in rows:
        enriched = deepcopy(row)
        publication = enriched["metadata"]["publication_time"]
        scope, interval, note = annotations[str(row["record_id"])]
        publication["working_scope"] = scope
        publication["boundary_interval"] = interval or None
        publication["boundary_note"] = note
        enriched_rows.append(enriched)
    _atomic_text(
        output_dir / "records.jsonl",
        "".join(
            json.dumps(row, ensure_ascii=False) + "\n" for row in enriched_rows
        ),
    )
    _write_csv(output_dir / "records.csv", csv_headers, csv_rows, annotations)
    _write_xlsx(source / "records.xlsx", output_dir / "records.xlsx", annotations)

    summary = {
        "source_run": str(source_run),
        "reference_date": "2026-07-21",
        "strict_cutoff": "2026-04-21",
        "working_policy": "include_three_month_boundary_for_recall",
        "total_records": len(rows),
        "strict_records": strict_count,
        "boundary_uncertain_records": boundary_count,
        "excluded_four_months_and_older": True,
        "boundary_interval": "2026-03-22 ~ 2026-04-21",
        "output_files": ["records.jsonl", "records.csv", "records.xlsx"],
    }
    _atomic_text(
        output_dir / "boundary-summary.json",
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n",
    )
    _atomic_text(
        output_dir / "边界说明.md",
        "# 实现网三个月任务工作清单\n\n"
        f"- 总记录：{len(rows)}\n"
        f"- 严格三个月内：{strict_count}\n"
        f"- ‘3个月前’边界记录：{boundary_count}\n"
        "- 边界范围：约 2026-03-22 至 2026-04-21\n"
        "- 使用方式：为避免漏单，边界记录纳入工作清单，但不作为严格三个月内记录。\n"
        "- ‘4个月前’及更早记录继续排除。\n",
    )
    return summary


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Create a Shixian working list with explicit boundary notes."
    )
    parser.add_argument("source_run", type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--expected-total", type=int)
    parser.add_argument("--expected-boundary", type=int)
    args = parser.parse_args()
    try:
        summary = finalize(
            args.source_run,
            args.output_dir,
            expected_total=args.expected_total,
            expected_boundary=args.expected_boundary,
        )
    except BoundaryNoteError as exc:
        print(f"ERROR: {exc}")
        return 1
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
