from __future__ import annotations

import argparse
from collections import Counter
import csv
from datetime import date
import json
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook


SECRET_PATTERN = re.compile(
    r"(?i)(?:HWWAFSESID|PHPSESSID|access_token|imwk_token|Authorization|Cookie)"
    r"\s*[:=]\s*[\" ]?([A-Za-z0-9%+/_=.-]{8,})"
)


def parse_publication_date(value: object) -> date:
    match = re.search(r"(20\d{2})[-./](\d{1,2})[-./](\d{1,2})", str(value))
    if match is None:
        raise ValueError(f"bad publication date: {value}")
    return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [
        json.loads(line)
        for line in path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]


DEFAULT_RUN_DIR = Path("outputs/外包--完整数据/一品威客")


def validate(project_root: Path, run_dir: Path | None = None) -> dict[str, Any]:
    selected_run_dir = run_dir or DEFAULT_RUN_DIR
    if not selected_run_dir.is_absolute():
        selected_run_dir = project_root / selected_run_dir
    base = selected_run_dir / "epwk"
    jsonl_rows = read_jsonl(base / "records.jsonl")
    jsonl_urls = [row.get("detail_url") for row in jsonl_rows]
    url_set = set(jsonl_urls)
    dates = [parse_publication_date(row.get("published_at_raw")) for row in jsonl_rows]
    cutoff = date(2026, 4, 21)
    as_of = date(2026, 7, 21)

    with (base / "records.csv").open(
        "r", encoding="utf-8-sig", newline=""
    ) as handle:
        csv_rows = list(csv.DictReader(handle))
    csv_urls = {row.get("detail_url") for row in csv_rows}

    workbook = load_workbook(base / "records.xlsx", read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    headers = [
        str(value) if value is not None else ""
        for value in next(iterator)
    ]
    xlsx_rows = [dict(zip(headers, values)) for values in iterator]
    workbook.close()
    xlsx_urls = {row.get("detail_url") for row in xlsx_rows}

    report = json.loads((base / "report.json").read_text(encoding="utf-8"))
    manifest = json.loads(
        (base / "unresolved-manifest.json").read_text(encoding="utf-8")
    )
    unresolved = manifest.get("unresolved", [])
    unresolved_urls = {row.get("detail_url") for row in unresolved}
    status_counts = Counter(row.get("status") for row in unresolved)

    supplemental_urls = {
        "https://task.epwk.com/1157358/",
        "https://task.epwk.com/1152909/",
    }
    failure_recovery_rows = [
        row
        for row in jsonl_rows
        if isinstance(row.get("metadata"), dict)
        and isinstance(row["metadata"].get("failure_recovery"), dict)
    ]
    supplemental_rows = [
        row for row in jsonl_rows if row.get("detail_url") in supplemental_urls
    ]
    mismatch_rows = [
        row
        for row in failure_recovery_rows
        if row.get("detail_url") not in supplemental_urls
    ]
    browser_rows = [
        row
        for row in jsonl_rows
        if isinstance(row.get("metadata"), dict)
        and isinstance(row["metadata"].get("browser_parse_gap_recovery"), dict)
    ]
    recovery_urls = {
        row.get("detail_url")
        for row in mismatch_rows + browser_rows + supplemental_rows
    }

    text_files = [
        base / "records.jsonl",
        base / "records.csv",
        base / "report.json",
        base / "report.md",
        base / "unresolved-manifest.json",
    ]
    secret_hits: dict[str, int] = {}
    for path in text_files:
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        matches = [
            value
            for value in SECRET_PATTERN.findall(text)
            if value.lower() not in {"redacted", "[redacted]"}
        ]
        if matches:
            secret_hits[path.name] = len(matches)

    summary: dict[str, Any] = {
        "jsonl_rows": len(jsonl_rows),
        "csv_rows": len(csv_rows),
        "xlsx_rows": len(xlsx_rows),
        "report_record_count": report.get("quality", {}).get("record_count"),
        "unique_urls": len(url_set),
        "csv_url_set_matches": csv_urls == url_set,
        "xlsx_url_set_matches": xlsx_urls == url_set,
        "core_complete": sum(
            row.get("core_complete") is True for row in jsonl_rows
        ),
        "date_min": min(dates).isoformat(),
        "date_max": max(dates).isoformat(),
        "outside_window": sum(
            not (cutoff <= value <= as_of) for value in dates
        ),
        "unresolved_count_field": manifest.get("unresolved_count"),
        "unresolved_entries": len(unresolved),
        "unresolved_unique_urls": len(unresolved_urls),
        "unresolved_status_counts": dict(status_counts),
        "records_unresolved_overlap": len(url_set & unresolved_urls),
        "mismatch_recovery_rows": len(mismatch_rows),
        "browser_recovery_rows": len(browser_rows),
        "non_sitemap_recovery_rows": len(supplemental_rows),
        "recovery_unique_urls": len(recovery_urls),
        "recovery_present_in_final": len(recovery_urls & url_set),
        "recovery_missing": sorted(recovery_urls - url_set),
        "filtered_out": report.get("filtered_out_count"),
        "source_coverage_inputs": len(report.get("source_coverage", {})),
        "accounting_total": (
            len(jsonl_rows)
            + report.get("filtered_out_count", 0)
            + len(unresolved)
        ),
        "manifest_accounting_total": manifest.get(
            "accounting", {}
        ).get("total_candidates"),
        "secret_hit_files": secret_hits,
    }
    checks = {
        "record_counts": (
            summary["jsonl_rows"]
            == summary["csv_rows"]
            == summary["xlsx_rows"]
            == summary["report_record_count"]
            == 2691
        ),
        "unique_and_exports": (
            summary["unique_urls"] == 2691
            and summary["csv_url_set_matches"]
            and summary["xlsx_url_set_matches"]
        ),
        "core_and_window": (
            summary["core_complete"] == 2691
            and summary["outside_window"] == 0
        ),
        "unresolved": (
            summary["unresolved_count_field"]
            == summary["unresolved_entries"]
            == summary["unresolved_unique_urls"]
            == 86
            and status_counts
            == Counter({"http_403_unresolved": 85, "parse_error": 1})
        ),
        "disjoint": summary["records_unresolved_overlap"] == 0,
        "recovery": (
            summary["mismatch_recovery_rows"] == 25
            and summary["browser_recovery_rows"] == 4
            and summary["non_sitemap_recovery_rows"] == 2
            and summary["recovery_unique_urls"]
            == summary["recovery_present_in_final"]
            == 31
            and not summary["recovery_missing"]
        ),
        "coverage_and_accounting": (
            summary["filtered_out"] == 1323
            and summary["source_coverage_inputs"] == 10
            and summary["accounting_total"]
            == summary["manifest_accounting_total"]
            == 4100
        ),
        "secret_scan": not summary["secret_hit_files"],
    }
    summary["checks"] = checks
    summary["all_passed"] = all(checks.values())
    return summary


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Validate the final merged EPWK artifacts."
    )
    parser.add_argument(
        "--project-root",
        type=Path,
        default=Path(__file__).resolve().parents[1],
    )
    parser.add_argument(
        "--run-dir",
        type=Path,
        default=DEFAULT_RUN_DIR,
        help="merged EPWK run directory, relative to project root by default",
    )
    args = parser.parse_args()
    summary = validate(args.project_root.resolve(), args.run_dir)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0 if summary["all_passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
