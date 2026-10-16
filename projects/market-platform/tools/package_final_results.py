from __future__ import annotations

import argparse
from decimal import Decimal
import json
import os
import re
from pathlib import Path
from typing import Any, Iterable
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from drv_crawler.categories import (
    CategoryParseError,
    validate_original_category_fields,
)


MAX_RECORDS_PER_SITE = 20_000
PRICE_SOURCE_PRIORITY = (
    "detail_summary",
    "list_page_embedded_state",
    "list_page",
    "platform_estimate",
)
BOUNDARY_SCOPE = "three_month_boundary_uncertain"
UNPRICED_MARKERS = {
    "待服务商报价",
    "带服务商报价",
    "待报价",
    "面议",
}


class PackageError(RuntimeError):
    pass


def _read_records(path: Path) -> list[dict[str, Any]]:
    try:
        rows = [
            json.loads(line)
            for line in path.read_text(encoding="utf-8").splitlines()
            if line.strip()
        ]
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise PackageError(f"cannot read {path}") from exc
    if not rows or len(rows) > MAX_RECORDS_PER_SITE:
        raise PackageError(f"invalid record count in {path}: {len(rows)}")
    record_ids = [str(row.get("record_id") or "") for row in rows]
    if not all(record_ids) or len(set(record_ids)) != len(record_ids):
        raise PackageError(f"missing or duplicate record_id in {path}")
    return rows


def _unique_texts(values: Iterable[object]) -> list[str]:
    seen: set[str] = set()
    output: list[str] = []
    for value in values:
        text = str(value or "").strip()
        if text and text not in seen:
            seen.add(text)
            output.append(text)
    return output


def _title(row: dict[str, Any]) -> str:
    value = str(row.get("detail_title_raw") or row.get("title_raw") or "").strip()
    if not value:
        raise PackageError(f"{row.get('record_id')} has no task/job title")
    return value


def _raw_price(row: dict[str, Any]) -> str:
    evidence = row.get("price_evidence")
    items = evidence if isinstance(evidence, list) else []
    for source in PRICE_SOURCE_PRIORITY:
        selected = _unique_texts(
            item.get("raw_text")
            for item in items
            if isinstance(item, dict) and item.get("price_source") == source
        )
        if selected:
            return " | ".join(selected)
    fallback = str(row.get("price_raw_joined") or "").strip()
    if fallback:
        return fallback
    raise PackageError(f"{row.get('record_id')} has no price")


def _decimal_number(value: Decimal) -> int | float:
    integral = value.to_integral_value()
    return int(integral) if value == integral else float(value)


def _clean_price(value: str) -> int | float | None:
    text = re.sub(r"\s+", "", value).replace("，", ",")
    text = text.replace("—", "-").replace("–", "-").replace("~", "-")
    text = text.replace("至", "-")
    text = text.replace(",", "")
    if text in UNPRICED_MARKERS:
        return None

    exact_match = re.fullmatch(
        r"(?:[￥¥])?(\d+(?:\.\d+)?)(万)?(?:元(?:/月)?)?", text
    )
    if exact_match:
        multiplier = Decimal(10_000) if exact_match.group(2) else Decimal(1)
        return _decimal_number(Decimal(exact_match.group(1)) * multiplier)

    monthly_match = re.fullmatch(r"(?:一般)?月薪(\d+(?:\.\d+)?)元", text)
    if monthly_match:
        return _decimal_number(Decimal(monthly_match.group(1)))

    upper_bound_match = re.fullmatch(
        r"(?:0-)?(\d+(?:\.\d+)?)元(?:以内)?", text
    )
    if upper_bound_match:
        return _decimal_number(Decimal(upper_bound_match.group(1)))

    yuan_range_match = re.fullmatch(
        r"(\d+(?:\.\d+)?)-(\d+(?:\.\d+)?)元", text
    )
    if yuan_range_match:
        low, high = map(Decimal, yuan_range_match.groups())
        if low > high:
            raise PackageError(f"price range is reversed: {value}")
        return _decimal_number((low + high) / 2)

    wan_range_match = re.fullmatch(
        r"(\d+(?:\.\d+)?)万-(\d+(?:\.\d+)?)万元", text
    )
    if wan_range_match:
        low, high = map(Decimal, wan_range_match.groups())
        if low > high:
            raise PackageError(f"price range is reversed: {value}")
        return _decimal_number((low + high) * Decimal(5_000))

    lower_bound_match = re.fullmatch(r"(\d+(?:\.\d+)?)万元以上", text)
    if lower_bound_match:
        return _decimal_number(Decimal(lower_bound_match.group(1)) * Decimal(10_000))
    raise PackageError(f"unsupported price format: {value}")


def _price(row: dict[str, Any]) -> int | float | None:
    return _clean_price(_raw_price(row))


def _has_real_price(row: dict[str, Any]) -> bool:
    return _price(row) is not None


def _category_levels(row: dict[str, Any]) -> tuple[str, str | None, str | None]:
    try:
        levels = validate_original_category_fields(
            str(row.get("site_name") or ""),
            row,
            detail_url=str(row.get("detail_url") or "") or None,
        )
    except CategoryParseError as exc:
        raise PackageError(
            f"{row.get('record_id')} has invalid original category: {exc}"
        ) from exc
    padded = (*levels, None, None)
    return padded[0], padded[1], padded[2]


def _is_boundary(row: dict[str, Any]) -> bool:
    metadata = row.get("metadata")
    publication = metadata.get("publication_time") if isinstance(metadata, dict) else None
    return bool(
        isinstance(publication, dict)
        and publication.get("working_scope") == BOUNDARY_SCOPE
    )


def _write_workbook(path: Path, rows: list[dict[str, Any]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "任务与价格"
    sheet.append(["任务/岗位", "分类一级", "分类二级", "分类三级", "价格"])
    for row in rows:
        level_1, level_2, level_3 = _category_levels(row)
        sheet.append([_title(row), level_1, level_2, level_3, _price(row)])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:E{sheet.max_row}"
    sheet.column_dimensions["A"].width = 60
    for column in "BCD":
        sheet.column_dimensions[column].width = 22
    sheet.column_dimensions["E"].width = 28
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    temporary = path.with_name(f".{path.name}.{uuid4().hex}.tmp")
    workbook.save(temporary)
    workbook.close()
    os.replace(temporary, path)


def _write_boundary_note(path: Path, strict_count: int, boundary_count: int) -> None:
    strict_start = 2
    strict_end = strict_count + 1
    boundary_start = strict_end + 1
    boundary_end = boundary_start + boundary_count - 1
    text = (
        "# 实现网边界说明\n\n"
        f"- 结果总数：{strict_count + boundary_count} 条。\n"
        f"- Excel 第 {strict_start}～{strict_end} 行：严格三个月内，共 {strict_count} 条。\n"
        f"- Excel 第 {boundary_start}～{boundary_end} 行：页面仅显示‘3个月前’的边界记录，共 {boundary_count} 条。\n"
        "- 边界记录时间约为 2026-03-22 至 2026-04-21；为避免漏单而纳入，但不作为严格三个月内记录。\n"
        "- 页面显示‘4个月前’及更早的记录未纳入。\n"
    )
    path.write_text(text, encoding="utf-8", newline="")


def _write_root_note(
    path: Path,
    *,
    counts: dict[str, int],
    strict_count: int,
    boundary_count: int,
    zbj_unpriced_filtered: int,
) -> None:
    text = (
        "# 任务与岗位数据说明\n\n"
        "本目录包含三个平台在本次采集范围内的任务或岗位、原站分类及数值化价格。\n\n"
        "## 结果汇总\n\n"
        f"- 猪八戒：{counts['猪八戒']} 条有明确预算的任务；另过滤 "
        f"{zbj_unpriced_filtered} 条‘待服务商报价’记录。\n"
        f"- 一品威客：{counts['一品威客']} 条任务。\n"
        f"- 实现网：{counts['实现网']} 条岗位，其中严格三个月内 "
        f"{strict_count} 条，‘3个月前’边界记录 {boundary_count} 条。\n\n"
        "## 分类口径\n\n"
        "- 猪八戒：使用每条记录所属的原站需求列表栏目，共六类。\n"
        "- 一品威客：使用任务详情页三级面包屑；其中 1 条详情已失效，使用"
        "搜索索引保存的原站栏目标签与当前原站分类体系恢复，并在完整数据元信息中保留证据。\n"
        "- 实现网：使用岗位详情页分类面包屑。\n"
        "- 标题只用于核对详情身份，不用于推断分类；标题与分类文字偶然相同不代表来源混用。\n\n"
        "## 价格口径\n\n"
        "价格列只保存 Excel 数值，不含货币和周期文字。精确价格直接取金额；"
        "区间取上下限中点；‘N元以内’取上限；‘N万元以上’取下限。"
        "实现网数值为原站月薪，另外两个平台为任务预算。\n"
    )
    path.write_text(text, encoding="utf-8", newline="")


def package_results(
    *,
    zbj_records: Path,
    epwk_records: Path,
    shixian_records: Path,
    output_dir: Path,
) -> dict[str, Any]:
    if output_dir.exists():
        raise PackageError(f"output already exists: {output_dir}")

    source_rows = {
        "猪八戒": _read_records(zbj_records),
        "一品威客": _read_records(epwk_records),
        "实现网": _read_records(shixian_records),
    }
    zbj_before = len(source_rows["猪八戒"])
    source_rows["猪八戒"] = [
        row for row in source_rows["猪八戒"] if _has_real_price(row)
    ]
    zbj_unpriced_filtered = zbj_before - len(source_rows["猪八戒"])
    shixian = source_rows["实现网"]
    strict_rows = [row for row in shixian if not _is_boundary(row)]
    boundary_rows = [row for row in shixian if _is_boundary(row)]
    if not boundary_rows:
        raise PackageError("Shixian source has no annotated boundary records")
    source_rows["实现网"] = strict_rows + boundary_rows

    temporary_root = output_dir.with_name(f".{output_dir.name}.{uuid4().hex}.tmp")
    temporary_root.mkdir(parents=True)
    for site_name, rows in source_rows.items():
        site_dir = temporary_root / site_name
        site_dir.mkdir()
        _write_workbook(site_dir / "任务与价格.xlsx", rows)
        if site_name == "实现网":
            _write_boundary_note(
                site_dir / "边界说明.md", len(strict_rows), len(boundary_rows)
            )
    _write_root_note(
        temporary_root / "说明.md",
        counts={site: len(rows) for site, rows in source_rows.items()},
        strict_count=len(strict_rows),
        boundary_count=len(boundary_rows),
        zbj_unpriced_filtered=zbj_unpriced_filtered,
    )
    os.replace(temporary_root, output_dir)
    return {
        "output_dir": str(output_dir),
        "counts": {site: len(rows) for site, rows in source_rows.items()},
        "shixian_strict": len(strict_rows),
        "shixian_boundary": len(boundary_rows),
        "zbj_unpriced_filtered": zbj_unpriced_filtered,
    }


def validate_package(output_dir: Path) -> dict[str, Any]:
    note_path = output_dir / "说明.md"
    if not note_path.is_file() or "分类口径" not in note_path.read_text(encoding="utf-8"):
        raise PackageError("missing or invalid root explanation")
    expected_files = {
        "猪八戒": {"任务与价格.xlsx"},
        "一品威客": {"任务与价格.xlsx"},
        "实现网": {"任务与价格.xlsx", "边界说明.md"},
    }
    counts: dict[str, int] = {}
    for site_name, names in expected_files.items():
        site_dir = output_dir / site_name
        actual = {path.name for path in site_dir.iterdir() if path.is_file()}
        allowed = names | {"分类价格统计.xlsx"}
        if not names.issubset(actual) or not actual.issubset(allowed):
            raise PackageError(
                f"unexpected files in {site_dir}: required {sorted(names)}, "
                f"allowed {sorted(allowed)}, got {sorted(actual)}"
            )
        workbook = load_workbook(site_dir / "任务与价格.xlsx", read_only=True)
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        rows = list(sheet.iter_rows(min_row=2))
        workbook.close()
        expected_headers = ["任务/岗位", "分类一级", "分类二级", "分类三级", "价格"]
        if headers != expected_headers:
            raise PackageError(f"unexpected headers for {site_name}: {headers}")
        if any(
            not row[0].value
            or not row[1].value
            or isinstance(row[4].value, bool)
            or not isinstance(row[4].value, (int, float))
            or row[4].value <= 0
            or row[4].data_type == "f"
            for row in rows
        ):
            raise PackageError(f"empty title/category/price in {site_name}")
        counts[site_name] = len(rows)
    return {"valid": True, "counts": counts}


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Package the three final crawler datasets with original categories."
    )
    parser.add_argument("--zbj-records", required=True, type=Path)
    parser.add_argument("--epwk-records", required=True, type=Path)
    parser.add_argument("--shixian-records", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    try:
        summary = package_results(
            zbj_records=args.zbj_records,
            epwk_records=args.epwk_records,
            shixian_records=args.shixian_records,
            output_dir=args.output_dir,
        )
        summary["validation"] = validate_package(args.output_dir)
    except (OSError, PackageError, ValueError) as exc:
        print(f"ERROR: {exc}")
        return 1
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
