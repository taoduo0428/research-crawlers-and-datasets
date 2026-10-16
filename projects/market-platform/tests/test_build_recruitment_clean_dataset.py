from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tools.build_recruitment_clean_dataset import (
    CLEAN_FIELDS,
    DatasetBuildError,
    build_site,
    clean_text,
    read_jsonl,
    _xlsx_cell,
)


def _write_jsonl(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows),
        encoding="utf-8",
    )


def _base_row(record_id: str, site_name: str, *, jd: str | None = None) -> dict[str, object]:
    return {
        "record_id": record_id,
        "site_name": site_name,
        "title_raw": "  Python  工程师 ",
        "description_raw": jd,
        "budget_raw": "20-30k",
        "category_level_1_raw": "IT/互联网",
        "category_level_2_raw": "后端开发",
        "category_level_3_raw": "Python",
        "detail_url": f"https://example.test/{record_id}",
        "metadata": {"keep": True},
        "price_evidence": [{"raw_text": "20-30k"}],
    }


def test_build_liepin_merges_jd_and_keeps_missing_rows_in_complete_metadata(tmp_path: Path) -> None:
    input_root = tmp_path / "input"
    detail_root = tmp_path / "details"
    complete_root = tmp_path / "complete"
    clean_root = tmp_path / "clean"
    _write_jsonl(
        input_root / "猎聘" / "岗位数据.jsonl",
        [_base_row("one", "liepin"), _base_row("two", "liepin")],
    )
    _write_jsonl(
        detail_root / "猎聘" / "详情补充.jsonl",
        [
            {
                "record_id": "one",
                "description_raw": "职责：<br>开发 API\r\n\r\n\r\n要求：本科",
                "detail_fetch_status": "ok",
                "detail_capture_method": "detail_html",
            }
        ],
    )

    stats = build_site("猎聘", input_root, detail_root, complete_root, clean_root)

    complete = list(read_jsonl(complete_root / "猎聘" / "岗位数据.jsonl"))
    cleaned = list(read_jsonl(clean_root / "猎聘" / "岗位数据.jsonl"))
    assert len(complete) == 2
    assert complete[0]["metadata"] == {"keep": True}
    assert complete[0]["description_raw"] == "职责：<br>开发 API\r\n\r\n\r\n要求：本科"
    assert complete[1]["description_raw"] == ""
    assert complete[1]["detail_fetch_status"] == "missing"
    assert cleaned == [
        {
            "职位": "Python 工程师",
            "JD": "职责：\n开发 API\n\n要求：本科",
            "薪资": "20-30K",
            "一级分类": "IT/互联网",
            "二级分类": "后端开发",
            "三级分类": "Python",
        }
    ]
    assert stats["complete_metadata"]["rows_without_jd"] == 1
    assert stats["cleaned"]["excluded_without_jd"] == 1

    with (clean_root / "猎聘" / "岗位数据.csv").open(
        "r", encoding="utf-8-sig", newline=""
    ) as handle:
        csv_rows = list(csv.DictReader(handle))
    assert tuple(csv_rows[0]) == CLEAN_FIELDS
    workbook = load_workbook(clean_root / "猎聘" / "岗位数据.xlsx", read_only=True)
    try:
        assert list(workbook.active.values)[0] == CLEAN_FIELDS
        assert sum(1 for _ in workbook.active.iter_rows(values_only=True)) == 2
    finally:
        workbook.close()


def test_build_51job_uses_existing_jd_and_preserves_complete_metadata(tmp_path: Path) -> None:
    input_root = tmp_path / "input"
    complete_root = tmp_path / "complete"
    clean_root = tmp_path / "clean"
    row = _base_row("job51-one", "51job", jd="岗位职责\n1. 数据分析")
    row["category_paths"] = [{"path_key": "A/B/C"}]
    _write_jsonl(input_root / "前程无忧" / "岗位数据.jsonl", [row])

    stats = build_site("前程无忧", input_root, tmp_path / "details", complete_root, clean_root)

    complete = next(read_jsonl(complete_root / "前程无忧" / "岗位数据.jsonl"))
    assert complete["category_paths"] == [{"path_key": "A/B/C"}]
    assert stats["cleaned"]["clean_rows"] == 1


def test_build_boss_uses_existing_jd_and_preserves_complete_metadata(tmp_path: Path) -> None:
    input_root = tmp_path / "input"
    complete_root = tmp_path / "complete"
    clean_root = tmp_path / "clean"
    row = _base_row("boss-one", "boss", jd="岗位职责\n1. 负责服务研发")
    row.update(
        {
            "category_path_raw": "互联网/AI > 后端开发 > Python",
            "category_source": "official_position_tree_query",
            "job_link": "https://www.zhipin.com/job_detail/boss-one.html",
        }
    )
    _write_jsonl(input_root / "BOSS" / "岗位数据.jsonl", [row])

    stats = build_site("BOSS", input_root, tmp_path / "details", complete_root, clean_root)

    complete = next(read_jsonl(complete_root / "BOSS" / "岗位数据.jsonl"))
    cleaned = next(read_jsonl(clean_root / "BOSS" / "岗位数据.jsonl"))
    assert complete == row
    assert cleaned == {
        "职位": "Python 工程师",
        "JD": "岗位职责\n1. 负责服务研发",
        "薪资": "20-30K",
        "一级分类": "IT/互联网",
        "二级分类": "后端开发",
        "三级分类": "Python",
    }
    assert stats["complete_metadata"]["rows_with_jd"] == 1
    assert stats["cleaned"]["clean_rows"] == 1


def test_build_site_supports_canonical_input_and_output_root(tmp_path: Path) -> None:
    complete_root = tmp_path / "complete"
    clean_root = tmp_path / "clean"
    row = _base_row("canonical", "51job", jd="岗位职责\n维护元数据")
    _write_jsonl(complete_root / "前程无忧" / "岗位数据.jsonl", [row])

    stats = build_site(
        "前程无忧",
        complete_root,
        tmp_path / "details",
        complete_root,
        clean_root,
    )

    completed = next(read_jsonl(complete_root / "前程无忧" / "岗位数据.jsonl"))
    assert completed == row
    assert stats["metadata_preservation_rows"] == 1
    assert stats["cleaned"]["clean_rows"] == 1

def test_duplicate_detail_prefers_nonempty_jd(tmp_path: Path) -> None:
    input_root = tmp_path / "input"
    detail_root = tmp_path / "details"
    _write_jsonl(input_root / "猎聘" / "岗位数据.jsonl", [_base_row("one", "liepin")])
    _write_jsonl(
        detail_root / "猎聘" / "详情补充.jsonl",
        [
            {"record_id": "one", "description_raw": "有效 JD"},
            {"record_id": "one", "description_raw": ""},
        ],
    )
    build_site("猎聘", input_root, detail_root, tmp_path / "complete", tmp_path / "clean")
    row = next(read_jsonl(tmp_path / "complete" / "猎聘" / "岗位数据.jsonl"))
    assert row["description_raw"] == "有效 JD"


def test_malformed_json_fails_without_publishing_dataset(tmp_path: Path) -> None:
    source = tmp_path / "input" / "前程无忧" / "岗位数据.jsonl"
    source.parent.mkdir(parents=True)
    source.write_text("{broken\n", encoding="utf-8")

    with pytest.raises(DatasetBuildError, match="不是合法 JSON"):
        build_site(
            "前程无忧",
            tmp_path / "input",
            tmp_path / "details",
            tmp_path / "complete",
            tmp_path / "clean",
        )
    assert not (tmp_path / "complete" / "前程无忧" / "岗位数据.jsonl").exists()


def test_clean_text_removes_html_controls_and_repeated_blanks() -> None:
    assert clean_text(" A\u200b  B ") == "A B"
    assert clean_text("职责<br>开发\n\n\n要求", multiline=True) == "职责\n开发\n\n要求"
    assert _xlsx_cell("第一行\r\n第二行") == ("第一行\n第二行", False)


def test_xlsx_keeps_equals_prefixed_jd_as_text(tmp_path: Path) -> None:
    input_root = tmp_path / "input"
    row = _base_row("formula-like", "51job", jd="=====岗位职责=====\n开发")
    _write_jsonl(input_root / "前程无忧" / "岗位数据.jsonl", [row])
    build_site(
        "前程无忧",
        input_root,
        tmp_path / "details",
        tmp_path / "complete",
        tmp_path / "clean",
    )

    workbook = load_workbook(
        tmp_path / "clean" / "前程无忧" / "岗位数据.xlsx",
        read_only=True,
        data_only=True,
    )
    try:
        values = list(workbook.active.values)
        assert values[1][1] == "=====岗位职责=====\n开发"
    finally:
        workbook.close()
