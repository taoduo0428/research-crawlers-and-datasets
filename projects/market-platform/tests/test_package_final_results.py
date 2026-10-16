from __future__ import annotations

import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook

from tools.package_final_results import _clean_price, package_results, validate_package


def make_record(
    record_id: str,
    *,
    site_name: str,
    title: str,
    price: str,
    source: str,
    boundary: bool = False,
) -> dict[str, object]:
    publication: dict[str, object] = {"status": "include"}
    if boundary:
        publication = {
            "status": "uncertain",
            "working_scope": "three_month_boundary_uncertain",
        }
    categories = {
        "zbj": {
            "category_level_1_raw": "软件开发",
            "category_level_2_raw": None,
            "category_level_3_raw": None,
            "category_path_raw": "软件开发",
            "category_source": "source_list_heading",
            "category_source_url": "https://www.zbj.com/xq/rjkf/r4/",
            "detail_url": "https://www.zbj.com/task/1",
        },
        "epwk": {
            "category_level_1_raw": "开发",
            "category_level_2_raw": "软件开发",
            "category_level_3_raw": "其他软件开发",
            "category_path_raw": "开发 > 软件开发 > 其他软件开发",
            "category_source": "detail_breadcrumb",
            "category_source_url": "https://task.epwk.com/1/",
            "detail_url": "https://task.epwk.com/1/",
        },
        "shixian": {
            "category_level_1_raw": "后端工程师",
            "category_level_2_raw": None,
            "category_level_3_raw": None,
            "category_path_raw": "后端工程师",
            "category_source": "detail_breadcrumb",
            "category_source_url": "https://shixian.com/job/all/backend",
            "detail_url": "https://shixian.com/jobs/1",
        },
    }
    return {
        "record_id": record_id,
        "site_name": site_name,
        "title_raw": title,
        "detail_title_raw": None,
        **categories[site_name],
        "category_identity_verified": True,
        "price_evidence": [{"raw_text": price, "price_source": source}],
        "price_raw_joined": price,
        "metadata": {"publication_time": publication},
    }


def write_jsonl(path: Path, rows: list[dict[str, object]]) -> None:
    path.write_text(
        "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows),
        encoding="utf-8",
    )


class PackageFinalResultsTests(unittest.TestCase):
    def test_packages_categories_prices_and_boundary_note(self) -> None:
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            zbj = root / "zbj.jsonl"
            epwk = root / "epwk.jsonl"
            shixian = root / "shixian.jsonl"
            write_jsonl(
                zbj,
                [
                    make_record(
                        "z1",
                        site_name="zbj",
                        title="猪八戒任务",
                        price="待服务商报价",
                        source="list_page_embedded_state",
                    ),
                    make_record(
                        "z2",
                        site_name="zbj",
                        title="有效价格任务",
                        price="0-500元",
                        source="list_page_embedded_state",
                    ),
                ],
            )
            write_jsonl(
                epwk,
                [
                    make_record(
                        "e1",
                        site_name="epwk",
                        title="一品任务",
                        price="1000元",
                        source="detail_summary",
                    )
                ],
            )
            write_jsonl(
                shixian,
                [
                    make_record(
                        "s2",
                        site_name="shixian",
                        title="边界任务",
                        price="2000元",
                        source="detail_summary",
                        boundary=True,
                    ),
                    make_record(
                        "s1",
                        site_name="shixian",
                        title="严格任务",
                        price="3000元",
                        source="detail_summary",
                    ),
                ],
            )
            output = root / "最终结果"

            summary = package_results(
                zbj_records=zbj,
                epwk_records=epwk,
                shixian_records=shixian,
                output_dir=output,
            )
            validation = validate_package(output)

            self.assertEqual(1, summary["shixian_strict"])
            self.assertEqual(1, summary["shixian_boundary"])
            self.assertEqual(1, summary["zbj_unpriced_filtered"])
            self.assertTrue(validation["valid"])
            zbj_workbook = load_workbook(
                output / "猪八戒" / "任务与价格.xlsx",
                read_only=True,
                data_only=True,
            )
            zbj_rows = list(zbj_workbook.active.iter_rows(values_only=True))
            zbj_workbook.close()
            self.assertEqual(
                [
                    ("任务/岗位", "分类一级", "分类二级", "分类三级", "价格"),
                    (
                        "有效价格任务",
                        "软件开发",
                        None,
                        None,
                        500,
                    ),
                ],
                zbj_rows,
            )
            workbook = load_workbook(
                output / "实现网" / "任务与价格.xlsx",
                read_only=True,
                data_only=True,
            )
            rows = list(workbook.active.iter_rows(values_only=True))
            workbook.close()
            self.assertEqual(
                ("任务/岗位", "分类一级", "分类二级", "分类三级", "价格"),
                rows[0],
            )
            self.assertEqual(
                ("严格任务", "后端工程师", None, None, 3000),
                rows[1],
            )
            self.assertEqual(
                ("边界任务", "后端工程师", None, None, 2000),
                rows[2],
            )
            self.assertIn(
                "Excel 第 3～3 行",
                (output / "实现网" / "边界说明.md").read_text(encoding="utf-8"),
            )
            self.assertIn(
                "标题只用于核对详情身份",
                (output / "说明.md").read_text(encoding="utf-8"),
            )

    def test_cleans_known_price_formats(self) -> None:
        self.assertEqual(500, _clean_price("￥ 500"))
        self.assertEqual(3000, _clean_price("一般月薪 3000 元"))
        self.assertEqual(500, _clean_price("0-500元"))
        self.assertEqual(15_000, _clean_price("1万至2万元"))
        self.assertEqual(1_000_000, _clean_price("100万元以上"))
        self.assertIsNone(_clean_price("待服务商报价"))


if __name__ == "__main__":
    unittest.main()
