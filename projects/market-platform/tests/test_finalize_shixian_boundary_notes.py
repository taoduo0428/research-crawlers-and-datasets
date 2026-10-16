from __future__ import annotations

import csv
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from tools.finalize_shixian_boundary_notes import (
    BOUNDARY_SCOPE,
    STRICT_SCOPE,
    BoundaryNoteError,
    finalize,
)


def record(record_id: str, *, boundary: bool) -> dict[str, object]:
    publication: dict[str, object] = {
        "status": "uncertain" if boundary else "include",
        "normalized_date": None if boundary else "2026-05-01",
        "interval_start": "2026-03-22" if boundary else None,
        "interval_end": "2026-04-21" if boundary else None,
    }
    if boundary:
        publication["boundary_assumption"] = "site_three_month_label"
    return {
        "record_id": record_id,
        "title_raw": record_id,
        "metadata": {"publication_time": publication},
    }


class FinalizeShixianBoundaryNotesTests(unittest.TestCase):
    def make_source(self, root: Path, rows: list[dict[str, object]]) -> Path:
        run_dir = root / "source"
        site_dir = run_dir / "shixian"
        site_dir.mkdir(parents=True)
        (site_dir / "records.jsonl").write_text(
            "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows),
            encoding="utf-8",
        )
        with (site_dir / "records.csv").open(
            "w", encoding="utf-8-sig", newline=""
        ) as handle:
            writer = csv.DictWriter(handle, fieldnames=["record_id", "title_raw"])
            writer.writeheader()
            for row in rows:
                writer.writerow(
                    {"record_id": row["record_id"], "title_raw": row["title_raw"]}
                )
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["record_id", "title_raw"])
        for row in rows:
            sheet.append([row["record_id"], row["title_raw"]])
        workbook.save(site_dir / "records.xlsx")
        workbook.close()
        return run_dir

    def test_writes_visible_boundary_columns_and_json_metadata(self) -> None:
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = self.make_source(
                root,
                [record("strict", boundary=False), record("boundary", boundary=True)],
            )
            output = root / "output"

            summary = finalize(
                source, output, expected_total=2, expected_boundary=1
            )

            self.assertEqual(1, summary["strict_records"])
            self.assertEqual(1, summary["boundary_uncertain_records"])
            json_rows = [
                json.loads(line)
                for line in (output / "records.jsonl").read_text(
                    encoding="utf-8"
                ).splitlines()
            ]
            self.assertEqual(
                STRICT_SCOPE,
                json_rows[0]["metadata"]["publication_time"]["working_scope"],
            )
            self.assertEqual(
                BOUNDARY_SCOPE,
                json_rows[1]["metadata"]["publication_time"]["working_scope"],
            )
            self.assertIn(
                "不作为严格三个月内记录",
                json_rows[1]["metadata"]["publication_time"]["boundary_note"],
            )
            with (output / "records.csv").open(
                encoding="utf-8-sig", newline=""
            ) as handle:
                csv_rows = list(csv.DictReader(handle))
            self.assertEqual(BOUNDARY_SCOPE, csv_rows[1]["publication_scope"])

            workbook = load_workbook(output / "records.xlsx", read_only=True)
            headers = [cell.value for cell in workbook.active[1]]
            workbook.close()
            self.assertIn("publication_boundary_note", headers)

    def test_rejects_boundary_record_without_interval(self) -> None:
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            bad = record("bad", boundary=True)
            bad["metadata"]["publication_time"]["interval_start"] = None  # type: ignore[index]
            source = self.make_source(root, [bad])

            with self.assertRaises(BoundaryNoteError):
                finalize(source, root / "output")


if __name__ == "__main__":
    unittest.main()
