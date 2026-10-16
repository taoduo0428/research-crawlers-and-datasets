from __future__ import annotations

from hashlib import sha256
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from tools.analyze_category_prices import (
    CategoryStatisticsError,
    _read_source_rows,
    generate_category_statistics,
)


HEADERS = ["任务/岗位", "分类一级", "分类二级", "分类三级", "价格"]


def write_source(path: Path, rows: list[tuple[object, ...]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "任务与价格"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def file_hash(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


class CategoryPriceStatisticsTests(unittest.TestCase):
    def test_generates_separate_workbooks_for_every_available_level(self) -> None:
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            write_source(
                root / "猪八戒" / "任务与价格.xlsx",
                [
                    ("任务1", "软件开发", None, None, 100),
                    ("任务2", "软件开发", None, None, 300),
                    ("任务3", "海报设计", None, None, 500),
                ],
            )
            write_source(
                root / "一品威客" / "任务与价格.xlsx",
                [
                    ("任务1", "开发", "软件开发", "小程序开发", 1000),
                    ("任务2", "开发", "软件开发", "小程序开发", 3000),
                    ("任务3", "开发", "网站开发", "其他网站", 500),
                    ("任务4", "营销", "软件开发", "小程序开发", 9000),
                ],
            )
            write_source(
                root / "实现网" / "任务与价格.xlsx",
                [
                    ("岗位1", "后端工程师", None, None, 1000),
                    ("岗位2", "后端工程师", None, None, 3000),
                ],
            )
            source_hashes = {
                site: file_hash(root / site / "任务与价格.xlsx")
                for site in ("猪八戒", "一品威客", "实现网")
            }

            summary = generate_category_statistics(root)

            self.assertEqual({"1": 2}, summary["猪八戒"]["levels"])
            self.assertEqual(
                {"1": 2, "2": 3, "3": 3}, summary["一品威客"]["levels"]
            )
            self.assertEqual({"1": 1}, summary["实现网"]["levels"])
            for site, expected_hash in source_hashes.items():
                self.assertEqual(
                    expected_hash, file_hash(root / site / "任务与价格.xlsx")
                )

            epwk = load_workbook(
                root / "一品威客" / "分类价格统计.xlsx",
                read_only=True,
                data_only=True,
            )
            self.assertEqual(
                ["一级分类统计", "二级分类统计", "三级分类统计"], epwk.sheetnames
            )
            level_one = list(epwk["一级分类统计"].iter_rows(values_only=True))
            level_two = list(epwk["二级分类统计"].iter_rows(values_only=True))
            epwk.close()
            self.assertEqual(("开发", 3, 1500, 1000), level_one[1])
            self.assertIn(("开发", "软件开发", 2, 2000, 2000), level_two)
            self.assertIn(("营销", "软件开发", 1, 9000, 9000), level_two)

            shixian = load_workbook(
                root / "实现网" / "分类价格统计.xlsx",
                read_only=True,
                data_only=True,
            )
            self.assertEqual(
                [("分类一级", "任务数量", "平均价格", "中位价格"),
                 ("后端工程师", 2, 2000, 2000)],
                list(shixian["一级分类统计"].iter_rows(values_only=True)),
            )
            shixian.close()

    def test_rejects_mixed_category_depths(self) -> None:
        with TemporaryDirectory() as temporary:
            path = Path(temporary) / "source.xlsx"
            write_source(
                path,
                [
                    ("任务1", "开发", "软件开发", "小程序开发", 1000),
                    ("任务2", "开发", None, None, 2000),
                ],
            )

            with self.assertRaises(CategoryStatisticsError):
                _read_source_rows(path)


if __name__ == "__main__":
    unittest.main()
