from __future__ import annotations

import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import tools.finalize_recruitment_results as finalizer_module
from tools.finalize_recruitment_results import (
    FinalizationError,
    finalize_recruitment_results,
)


SITES = ("boss", "job51", "liepin")


def _salary_evidence(value: str = "10-20K") -> list[dict[str, object]]:
    return [{"raw_text": value, "negotiable": False}]


def _base_record(site: str, index: int) -> dict[str, object]:
    return {
        "record_id": f"{site}-{index}",
        "site_name": "51job" if site == "job51" else site,
        "title_raw": f"Engineer {index}",
        "price_evidence": _salary_evidence(),
        "budget_raw": "10-20K",
        "publisher_name_raw": "Example Company",
        "location_raw": "Shanghai",
        "detail_url": f"https://example.test/{site}/{index}",
        "category_identity_verified": True,
        "metadata": {"salary_min": 10_000, "salary_max": 20_000},
    }


def _boss_paths() -> list[dict[str, str]]:
    return [
        {
            "category_level_1_code": "1000000",
            "category_level_1_raw": "Technology",
            "category_level_2_code": "1001000",
            "category_level_2_raw": "Backend",
            "category_level_3_code": "100101",
            "category_level_3_raw": "Python",
            "category_path_raw": "Technology > Backend > Python",
        },
        {
            "category_level_1_code": "1000000",
            "category_level_1_raw": "Technology",
            "category_level_2_code": "1002000",
            "category_level_2_raw": "Software",
            "category_level_3_code": "100101",
            "category_level_3_raw": "Python",
            "category_path_raw": "Technology > Software > Python",
        },
    ]


def _job51_paths() -> list[dict[str, str]]:
    return [
        {
            "level_1_code": "9F01",
            "level_1_name": "Technology",
            "level_2_code": "0100",
            "level_2_name": "Backend",
            "level_3_code": "0101",
            "level_3_name": "Python",
            "path_key": "9F01/0100/0101",
        },
        {
            "level_1_code": "9F01",
            "level_1_name": "Technology",
            "level_2_code": "0200",
            "level_2_name": "Software",
            "level_3_code": "0101",
            "level_3_name": "Python",
            "path_key": "9F01/0200/0101",
        },
    ]


def _boss_coverage(*, pending: int = 0, target: int = 50) -> dict[str, object]:
    paths = _boss_paths()
    return {
        "run_config": {"per_position_limit": target, "numeric_price_only": True},
        "position_codes_total": 1,
        "position_codes_completed": 1 - pending,
        "position_codes_pending": pending,
        "position_codes_deferred": 0,
        "position_codes_failed": 0,
        "quota_summary": {
            "quota_target": target,
            "quota_met": 1 - pending,
            "exhausted_below_quota": 0,
            "empty": 0,
            "page_limit_below_quota": 0,
            "failed": 0,
        },
        "positions": [
            {
                "position_code": "100101",
                "position_name": "Python",
                "category_path_raw": paths[0]["category_path_raw"],
                "category_paths_raw": [row["category_path_raw"] for row in paths],
                "status": "pending" if pending else "completed",
                "quota_target": target,
                "quota_status": "pending" if pending else "quota_met",
                "records_fetched": 0 if pending else target,
            }
        ],
    }


def _job51_coverage(*, pending: int = 0, target: int = 50) -> dict[str, object]:
    rows = []
    for path in _job51_paths():
        rows.append(
            {
                **path,
                "search_code": "0101",
                "status": "pending" if pending else "completed",
                "items_count": 0 if pending else target,
                "quota_outcome": "pending" if pending else "quota_met",
            }
        )
    return {
        "run_scope": {
            "per_category_limit": target,
            "numeric_salary_only": True,
        },
        "official_category_paths": 2,
        "unique_search_codes": 1,
        "quota_counts": {
            "unique_search_codes_total": 1,
            "target_per_unique_search_code": target,
            "numeric_salary_only": True,
            "quota_met": 1 - pending,
            "exhausted_below_quota": 0,
            "empty": 0,
            "failed": 0,
            "pending": pending,
            "page_limit_below_quota": 0,
            "running": 0,
        },
        "categories": rows,
    }


def _liepin_coverage(*, pending: int = 0, target: int = 50) -> dict[str, object]:
    return {
        "quota_target": target,
        "numeric_salary_only": True,
        "quota_outcomes": {
            "quota_met": 1 - pending,
            "exhausted_below_quota": 0,
            "empty": 0,
            "failed": 0,
            "pending": pending,
            "page_limit_below_quota": 0,
            "running": 0,
        },
        "summary": {
            "categories_total": 1,
            "categories_done": 1 - pending,
            "categories_empty": 0,
            "categories_failed": 0,
            "categories_pending": pending,
            "categories_running": 0,
            "records_total": 0 if pending else target,
        },
        "categories": [
            {
                "url": "https://www.liepin.com/career/python/",
                "path": "Technology > Backend > Python",
                "level_1": "Technology",
                "level_2": "Backend",
                "level_3": "Python",
                "status": "pending" if pending else "done",
                "quota_target": target,
                "quota_outcome": "pending" if pending else "quota_met",
                "record_count": 0 if pending else target,
            }
        ],
    }


def _write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


def _write_jsonl(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows),
        encoding="utf-8",
    )


def _write_scope(root: Path, site: str, *, pending: int = 0) -> Path:
    scope = root / site
    coverage = {
        "boss": _boss_coverage,
        "job51": _job51_coverage,
        "liepin": _liepin_coverage,
    }[site](pending=pending)
    _write_json(scope / "coverage.json", coverage)

    if pending:
        rows: list[dict[str, object]] = []
    elif site == "boss":
        paths = _boss_paths()
        rows = []
        for index in range(50):
            record = _base_record(site, index)
            record.update(
                {
                    "position_code": "100101",
                    "category_level_1_raw": paths[0]["category_level_1_raw"],
                    "category_level_2_raw": paths[0]["category_level_2_raw"],
                    "category_level_3_raw": paths[0]["category_level_3_raw"],
                    "category_path_raw": paths[0]["category_path_raw"],
                    "category_matches": [
                        {
                            "position_code": "100101",
                            "position_name": "Python",
                            "category_paths": paths,
                            "category_paths_raw": [
                                path["category_path_raw"] for path in paths
                            ],
                        }
                    ],
                }
            )
            rows.append(record)
        _write_json(scope / "jobs.json", {"jobs": rows})
        return scope
    elif site == "job51":
        paths = _job51_paths()
        rows = []
        for index in range(50):
            record = _base_record(site, index)
            record.update(
                {
                    "category_level_1_raw": paths[0]["level_1_name"],
                    "category_level_2_raw": paths[0]["level_2_name"],
                    "category_level_3_raw": paths[0]["level_3_name"],
                    "category_path_raw": "Technology > Backend > Python",
                    "category_paths": paths,
                }
            )
            record["metadata"] = {
                **record["metadata"],
                "source_function_code": "0101",
                "category_paths": paths,
            }
            rows.append(record)
    else:
        rows = []
        for index in range(50):
            record = _base_record(site, index)
            record.update(
                {
                    "category_level_1_raw": "Technology",
                    "category_level_2_raw": "Backend",
                    "category_level_3_raw": "Python",
                    "category_path_raw": "Technology > Backend > Python",
                    "category_source_url": "https://www.liepin.com/career/python/",
                    "list_url": "https://www.liepin.com/career/python/",
                }
            )
            rows.append(record)
    _write_jsonl(scope / "jobs.jsonl", rows)
    return scope


def _write_probe_scope(root: Path, site: str) -> Path:
    scope = root / site
    coverage = {
        "boss": _boss_coverage,
        "job51": _job51_coverage,
        "liepin": _liepin_coverage,
    }[site](target=1)
    if site == "liepin":
        # The persisted Liepin per1 probe predates explicit quota metadata.
        coverage.pop("quota_target")
        coverage.pop("quota_outcomes")
        for row in coverage["categories"]:
            row.pop("quota_target")
    _write_json(scope / "coverage.json", coverage)
    return scope


class FinalizeRecruitmentResultsTests(unittest.TestCase):
    def _fixtures(
        self, root: Path
    ) -> tuple[dict[str, Path], dict[str, Path]]:
        per50 = {
            site: _write_scope(root / "per50", site) for site in SITES
        }
        per1 = {
            site: _write_probe_scope(root / "per1", site) for site in SITES
        }
        return per50, per1

    def test_finalizes_three_sites_and_expands_alias_relations(self) -> None:
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            per50, per1 = self._fixtures(root)
            output = root / "final"

            report = finalize_recruitment_results(
                per50_scopes=per50,
                per1_scopes=per1,
                output_dir=output,
            )

            self.assertTrue(report["valid"])
            self.assertEqual(150, report["totals"]["records"])
            self.assertEqual(150, report["totals"]["category_associations"])
            self.assertEqual(250, report["totals"]["expanded_path_relations"])
            self.assertTrue((output / "validation.json").is_file())
            for site, expected_relations in (
                ("boss", 100),
                ("job51", 100),
                ("liepin", 50),
            ):
                site_report = report["sites"][site]
                self.assertTrue(site_report["inventory_matches_probe"])
                self.assertIsNotNone(site_report["probe_per_category_validation"])
                self.assertEqual(50, site_report["records"])
                self.assertEqual(expected_relations, site_report["expanded_path_relations"])
                workbook = load_workbook(
                    output / site / "jobs.xlsx", read_only=True, data_only=True
                )
                self.assertEqual(51, sum(1 for _ in workbook.active.iter_rows()))
                workbook.close()
                workbook = load_workbook(
                    output / site / "job_category_relations.xlsx",
                    read_only=True,
                    data_only=True,
                )
                self.assertEqual(
                    expected_relations + 1,
                    sum(1 for _ in workbook.active.iter_rows()),
                )
                workbook.close()
            self.assertEqual(
                "categories[].record_count<=1 (legacy coverage)",
                report["sites"]["liepin"]["probe_per_category_validation"],
            )

    def test_job51_record_requires_every_official_alias(self) -> None:
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            per50, per1 = self._fixtures(root)
            path = per50["job51"] / "jobs.jsonl"
            rows = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()]
            for row in rows:
                row["category_paths"] = row["category_paths"][:1]
                row["metadata"]["category_paths"] = row["category_paths"]
            _write_jsonl(path, rows)
            output = root / "final"

            with self.assertRaisesRegex(FinalizationError, "official category aliases"):
                finalize_recruitment_results(
                    per50_scopes=per50,
                    per1_scopes=per1,
                    output_dir=output,
                )

            self.assertFalse(output.exists())

    def test_boss_record_path_fields_must_match_inventory(self) -> None:
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            per50, per1 = self._fixtures(root)
            path = per50["boss"] / "jobs.json"
            payload = json.loads(path.read_text(encoding="utf-8"))
            for row in payload["jobs"]:
                row["category_matches"][0]["category_paths"][0][
                    "category_level_2_raw"
                ] = "Changed Backend"
            _write_json(path, payload)
            output = root / "final"

            with self.assertRaisesRegex(FinalizationError, "path fields differ"):
                finalize_recruitment_results(
                    per50_scopes=per50,
                    per1_scopes=per1,
                    output_dir=output,
                )

            self.assertFalse(output.exists())

    def test_quota_counts_must_match_unique_category_outcomes(self) -> None:
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            per50, per1 = self._fixtures(root)
            coverage_path = per50["job51"] / "coverage.json"
            coverage = json.loads(coverage_path.read_text(encoding="utf-8"))
            coverage["quota_counts"]["quota_met"] = 0
            coverage["quota_counts"]["exhausted_below_quota"] = 1
            _write_json(coverage_path, coverage)
            output = root / "final"

            with self.assertRaisesRegex(FinalizationError, "quota counts"):
                finalize_recruitment_results(
                    per50_scopes=per50,
                    per1_scopes=per1,
                    output_dir=output,
                )

            self.assertFalse(output.exists())

    def test_probe_scope_must_be_per_category_one(self) -> None:
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            per50, per1 = self._fixtures(root)
            coverage_path = per1["job51"] / "coverage.json"
            coverage = json.loads(coverage_path.read_text(encoding="utf-8"))
            coverage["run_scope"]["per_category_limit"] = 50
            coverage["quota_counts"]["target_per_unique_search_code"] = 50
            _write_json(coverage_path, coverage)
            output = root / "final"

            with self.assertRaisesRegex(FinalizationError, "per-category target"):
                finalize_recruitment_results(
                    per50_scopes=per50,
                    per1_scopes=per1,
                    output_dir=output,
                )

            self.assertFalse(output.exists())

    def test_saved_workbook_row_count_is_verified_before_publish(self) -> None:
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            per50, per1 = self._fixtures(root)
            output = root / "final"
            original_save = finalizer_module._save_workbook

            def save_then_truncate(workbook: Workbook, path: Path) -> None:
                original_save(workbook, path)
                if path.parent.name == "boss" and path.name == "jobs.xlsx":
                    truncated = Workbook()
                    truncated.active.title = "岗位"
                    truncated.active.append(["记录ID"])
                    truncated.save(path)
                    truncated.close()

            with patch.object(
                finalizer_module,
                "_save_workbook",
                side_effect=save_then_truncate,
            ):
                with self.assertRaisesRegex(FinalizationError, "workbook row count"):
                    finalize_recruitment_results(
                        per50_scopes=per50,
                        per1_scopes=per1,
                        output_dir=output,
                    )

            self.assertFalse(output.exists())
            self.assertEqual([], list(root.glob(".final.tmp-*")))

    def test_inventory_drift_rejects_atomic_publish(self) -> None:
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            per50, per1 = self._fixtures(root)
            coverage_path = per1["job51"] / "coverage.json"
            coverage = json.loads(coverage_path.read_text(encoding="utf-8"))
            coverage["categories"][0]["level_2_code"] = "9999"
            coverage["categories"][0]["path_key"] = "9F01/9999/0101"
            _write_json(coverage_path, coverage)
            output = root / "final"

            with self.assertRaisesRegex(FinalizationError, "inventory"):
                finalize_recruitment_results(
                    per50_scopes=per50,
                    per1_scopes=per1,
                    output_dir=output,
                )

            self.assertFalse(output.exists())
            self.assertEqual([], list(root.glob(".final.tmp-*")))

    def test_declared_inventory_total_mismatch_rejects_publish(self) -> None:
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            per50, per1 = self._fixtures(root)
            coverage_path = per50["liepin"] / "coverage.json"
            coverage = json.loads(coverage_path.read_text(encoding="utf-8"))
            coverage["summary"]["categories_total"] = 2
            _write_json(coverage_path, coverage)
            output = root / "final"

            with self.assertRaisesRegex(FinalizationError, "inventory total"):
                finalize_recruitment_results(
                    per50_scopes=per50,
                    per1_scopes=per1,
                    output_dir=output,
                )

            self.assertFalse(output.exists())

    def test_inventory_name_drift_rejects_publish(self) -> None:
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            per50, per1 = self._fixtures(root)
            coverage_path = per1["job51"] / "coverage.json"
            coverage = json.loads(coverage_path.read_text(encoding="utf-8"))
            coverage["categories"][0]["level_3_name"] = "Changed Name"
            _write_json(coverage_path, coverage)
            output = root / "final"

            with self.assertRaisesRegex(FinalizationError, "inventory"):
                finalize_recruitment_results(
                    per50_scopes=per50,
                    per1_scopes=per1,
                    output_dir=output,
                )

            self.assertFalse(output.exists())

    def test_status_count_conservation_rejects_publish(self) -> None:
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            per50, per1 = self._fixtures(root)
            coverage_path = per50["boss"] / "coverage.json"
            coverage = json.loads(coverage_path.read_text(encoding="utf-8"))
            coverage["position_codes_completed"] = 0
            _write_json(coverage_path, coverage)
            output = root / "final"

            with self.assertRaisesRegex(FinalizationError, "status total"):
                finalize_recruitment_results(
                    per50_scopes=per50,
                    per1_scopes=per1,
                    output_dir=output,
                )

            self.assertFalse(output.exists())

    def test_unfinished_quota_rejects_publish(self) -> None:
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            per50, per1 = self._fixtures(root)
            _write_scope(root / "per50", "liepin", pending=1)
            output = root / "final"

            with self.assertRaisesRegex(FinalizationError, "pending"):
                finalize_recruitment_results(
                    per50_scopes=per50,
                    per1_scopes=per1,
                    output_dir=output,
                )

            self.assertFalse(output.exists())

    def test_category_above_fifty_rejects_publish(self) -> None:
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            per50, per1 = self._fixtures(root)
            coverage_path = per50["job51"] / "coverage.json"
            coverage = json.loads(coverage_path.read_text(encoding="utf-8"))
            for row in coverage["categories"]:
                row["items_count"] = 51
            coverage["record_count"] = 51
            _write_json(coverage_path, coverage)
            output = root / "final"

            with self.assertRaisesRegex(FinalizationError, "maximum is 50"):
                finalize_recruitment_results(
                    per50_scopes=per50,
                    per1_scopes=per1,
                    output_dir=output,
                )

            self.assertFalse(output.exists())

    def test_non_numeric_salary_rejects_publish(self) -> None:
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            per50, per1 = self._fixtures(root)
            path = per50["liepin"] / "jobs.jsonl"
            rows = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()]
            rows[0]["price_evidence"] = _salary_evidence("面议")
            rows[0]["budget_raw"] = "面议"
            rows[0]["metadata"] = {}
            _write_jsonl(path, rows)
            output = root / "final"

            with self.assertRaisesRegex(FinalizationError, "numeric salary"):
                finalize_recruitment_results(
                    per50_scopes=per50,
                    per1_scopes=per1,
                    output_dir=output,
                )

            self.assertFalse(output.exists())

    def test_duplicate_category_relation_rejects_publish(self) -> None:
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            per50, per1 = self._fixtures(root)
            path = per50["job51"] / "jobs.jsonl"
            rows = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()]
            rows[0]["category_paths"].append(dict(rows[0]["category_paths"][0]))
            rows[0]["metadata"]["category_paths"] = rows[0]["category_paths"]
            _write_jsonl(path, rows)
            output = root / "final"

            with self.assertRaisesRegex(FinalizationError, "duplicate category relation"):
                finalize_recruitment_results(
                    per50_scopes=per50,
                    per1_scopes=per1,
                    output_dir=output,
                )

            self.assertFalse(output.exists())

    def test_existing_output_is_never_overwritten(self) -> None:
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            per50, per1 = self._fixtures(root)
            output = root / "final"
            output.mkdir()
            marker = output / "keep.txt"
            marker.write_text("keep", encoding="utf-8")

            with self.assertRaisesRegex(FinalizationError, "already exists"):
                finalize_recruitment_results(
                    per50_scopes=per50,
                    per1_scopes=per1,
                    output_dir=output,
                )

            self.assertEqual("keep", marker.read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
